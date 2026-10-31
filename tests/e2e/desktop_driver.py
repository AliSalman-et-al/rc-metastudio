# SPDX-FileCopyrightText: 2026 Ali Salman and RC MetaStudio contributors
# SPDX-License-Identifier: GPL-3.0-or-later
"""Drive application controls with the production R backend."""

import json
import time
import cProfile
import os
import shutil
import sys
import traceback
from typing import Callable, cast
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import patch

_startup_profile = cProfile.Profile() if os.environ.get("RCMS_PROFILE_STARTUP") else None
_startup_phases = {"driver:entry": time.time()}
if _startup_profile is not None:
    _startup_profile.enable()

from PyQt6 import QtCore, QtGui, QtWidgets
from PyQt6.QtSvg import QSvgRenderer
from PyQt6.QtTest import QTest

from tests.python.gui.test_types import key_click, mouse_click, required

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(sys.argv[2])

# Keep preferences and working artifacts out of the user's normal workspace.
QtCore.QSettings.setDefaultFormat(QtCore.QSettings.Format.IniFormat)
QtCore.QSettings.setPath(
    QtCore.QSettings.Format.IniFormat, QtCore.QSettings.Scope.UserScope, str(WORK)
)
QtCore.QStandardPaths.setTestModeEnabled(True)

from rc_metastudio import r_backend



from rc_metastudio.qt6_ui import prepare_generated_ui_imports

prepare_generated_ui_imports()

from rc_metastudio import analysis_setup_dialog, automation, results_window, settings
from tests.analysis_regression.golden.support.golden_analysis import (
    parsed_numeric_sections,
)

EVIDENCE = (
    Path(os.environ["RCMS_E2E_EVIDENCE_DIR"])
    if os.environ.get("RCMS_E2E_EVIDENCE_DIR")
    else None
)
STEP = 0


def record_step(action, *widgets, **observations):
    """Capture the action and its observed state while the relevant UI still exists."""
    global STEP
    if EVIDENCE is None:
        return
    STEP += 1
    captures = []
    from tests.e2e.display_labels import visible_label_inventory, plot_label_inventory
    groups = window.model.arm_names if window.model is not None else []
    user_names = [name for name in groups if name not in {"Group", "Intervention", "Control"}]
    if window.model is not None:
        user_names += [s.name for s in window.model.dataset.sheets]
        user_names += [c["name"] for c in window.model.current_sheet.covariates] if window.model.current_sheet else []
        user_names += [study.label for study in window.model.get_studies(False)]
    observations["visible_labels"] = [entry for widget in widgets for entry in visible_label_inventory(widget, user_names)]
    plot_paths = [value for value in observations.get("plots", {}).values() if isinstance(value, str)]
    for widget in widgets:
        result = getattr(widget, "results", None)
        if hasattr(result, "display_images"):
            plot_paths.extend(result.display_images.values())
    observations["visible_plot_labels"] = plot_label_inventory(plot_paths, user_names)
    for index, widget in enumerate(widgets):
        path = EVIDENCE / f"{STEP:02d}-{index}-{action}.png"
        if app.platformName() in {"windows", "xcb", "cocoa"}:
            widget.raise_()
            widget.activateWindow()
            app.processEvents()
            cast(Callable[[int], None], QTest.qWait)(400)
            if observations.get("capture_surface") == "desktop-composite":
                capture = widget.screen().grabWindow(0)
                ratio = capture.devicePixelRatio()
                geometry = widget.frameGeometry()
                origin = widget.screen().geometry().topLeft()
                capture = capture.copy(round((geometry.x() - origin.x()) * ratio),
                                       round((geometry.y() - origin.y()) * ratio),
                                       round(geometry.width() * ratio), round(geometry.height() * ratio))
            else:
                capture = widget.screen().grabWindow(int(widget.winId()))
        else:
            capture = widget.grab()
        assert capture.save(str(path), "PNG"), path
        captures.append(path.name)
        if app.platformName() in {"windows", "xcb", "cocoa"} and action in {"analysis-results", "matrix-results", "small-study-results", "family-plot-cancelled"}:
            report_checked = widget.report_action.isChecked()
            original_nav_item = widget.nav_tree.currentItem()
            widget.report_action.setChecked(True)
            app.processEvents()
            cast(Callable[[int], None], QTest.qWait)(400)
            report_path = EVIDENCE / f"{STEP:02d}-{index}-{action}-report.png"
            assert widget.screen().grabWindow(int(widget.winId())).save(str(report_path), "PNG")
            captures.append(report_path.name)
            (EVIDENCE / f"{STEP:02d}-report.txt").write_text(widget.report_text.toPlainText(), encoding="utf-8")
            if widget.tables_action.isEnabled():
                widget.tables_action.trigger()
                app.processEvents()
                cast(Callable[[int], None], QTest.qWait)(400)
                table_path = EVIDENCE / f"{STEP:02d}-{index}-{action}-table.png"
                assert widget.screen().grabWindow(int(widget.winId())).save(str(table_path), "PNG")
                captures.append(table_path.name)
                for table_index, table in enumerate(widget._report_tables):
                    if table.title == "Test of subgroup differences":
                        widget.table_selector.setCurrentIndex(table_index)
                        app.processEvents()
                        cast(Callable[[int], None], QTest.qWait)(400)
                        subgroup_path = EVIDENCE / f"{STEP:02d}-{index}-{action}-subgroup-test.png"
                        assert widget.screen().grabWindow(int(widget.winId())).save(str(subgroup_path), "PNG")
                        captures.append(subgroup_path.name)
                        widget.table_selector.setCurrentIndex(0)
            widget.report_action.setChecked(report_checked)
            widget._show_report(report_checked)
            app.processEvents()
            widget.nav_tree.setCurrentItem(getattr(widget, "_last_plot_item", original_nav_item))
            widget.report_action.setChecked(False)
            widget._show_report(False)
            scroll = widget.graphics_view.verticalScrollBar()
            position = scroll.value()
            widget._scroll_to_section(widget.nav_tree.currentItem())
            app.processEvents()
            cast(Callable[[int], None], QTest.qWait)(400)
            figure_path = EVIDENCE / f"{STEP:02d}-{index}-{action}-figure.png"
            assert widget.screen().grabWindow(int(widget.winId())).save(str(figure_path), "PNG")
            captures.append(figure_path.name)
            selected_plot = widget._selected_plot()
            if selected_plot is not None:
                bottom = selected_plot[1].sceneBoundingRect().bottom() * widget.graphics_view.transform().m22()
                scroll.setValue(round(bottom - widget.graphics_view.viewport().height() + 30))
            app.processEvents()
            cast(Callable[[int], None], QTest.qWait)(400)
            bottom_path = EVIDENCE / f"{STEP:02d}-{index}-{action}-bottom.png"
            assert widget.screen().grabWindow(int(widget.winId())).save(str(bottom_path), "PNG")
            captures.append(bottom_path.name)
            scroll.setValue(position)
            widget.nav_tree.setCurrentItem(original_nav_item)
            widget.report_action.setChecked(report_checked)
            widget._show_report(report_checked)
            app.processEvents()
    artifacts = {}
    for index, (title, source) in enumerate(observations.get("plots", {}).items()):
        source = Path(source)
        destination = EVIDENCE / f"{STEP:02d}-{index}-{source.name}"
        shutil.copy2(source, destination)
        artifacts[title] = destination.name
    event = {
        "step": STEP,
        "action": action,
        "utc": datetime.now(UTC).isoformat(),
        "captures": captures,
        "artifacts": artifacts,
        **observations,
    }
    with (EVIDENCE / "steps.jsonl").open("a", encoding="utf-8") as stream:
        stream.write(json.dumps(event, ensure_ascii=False) + "\n")


def doctor(app, window):
    from importlib import metadata

    from rc_metastudio import r_bridge
    from scripts.r_dependency_policy import load_policy

    policy = load_policy(ROOT / "config/r-dependencies.json")
    r_version = r_bridge.get_r_version_string()
    assert r_version.startswith("R version " + policy["r_version"] + " "), r_version
    assert (
        metadata.version("rpy2")
        == json.loads((ROOT / "config/r-dependencies.json").read_text())[
            "target_runtime"
        ]["rpy2"]
    )
    # packageVersion canonicalizes hyphens to dots; the manifest uses DCF spellings.
    versions = {
        name: str(
            r_bridge.execute_r_function("packageDescription", name).rx2("Version")[0]
        )
        for name in policy["pinned_authorities"]
    }
    assert versions == policy["pinned_authorities"], versions
    description = (ROOT / "r/RCMetaR/DESCRIPTION").read_text()
    rcmetar_version = next(
        line.split(":", 1)[1].strip()
        for line in description.splitlines()
        if line.startswith("Version:")
    )
    assert r_bridge.get_r_package_version("RCMetaR") == rcmetar_version
    assert window.isVisible()
    assert Path(settings.get_base_path()).resolve() == WORK.resolve()
    font_families = QtGui.QFontDatabase.families()
    font = QtGui.QRawFont.fromFont(app.font())
    assert font_families and font.isValid(), ("Qt has no usable fonts", app.font().toString())
    assert all(font.supportsCharacter(ord(char)) for char in "Aa012τ…"), (
        "Qt's default font lacks required glyphs", font.familyName(), app.font().toString()
    )
    identity = {
        "r": r_version,
        "r_libraries": r_bridge.get_r_library_paths(),
        "r_home": os.environ.get("R_HOME"),
        "authorities": versions,
        "rcmetar": rcmetar_version,
        "rpy2": metadata.version("rpy2"),
        "qt": QtCore.qVersion(),
        "pyqt": QtCore.PYQT_VERSION_STR,
        "platform_plugin": app.platformName(),
        "font": {
            "family": font.familyName(),
            "style": font.styleName(),
            "family_count": len(font_families),
            "required_glyphs": "Aa012τ…",
            "directory": os.environ.get("QT_QPA_FONTDIR"),
        },
        "application": app.applicationName(),
        "qt_build_root": os.environ.get("RCMS_QT6_BUILD_ROOT"),
        "application_base_path": settings.get_base_path(),
    }
    if EVIDENCE is not None:
        (EVIDENCE / "runtime.json").write_text(
            json.dumps(identity, indent=2) + "\n", encoding="utf-8"
        )
    record_step("ready", window, runtime=identity)


def analyze(app, window, method):
    from rc_metastudio.diagnostic_metrics_dialog import DiagnosticMetricsDialog

    assert window.action_go.isEnabled()
    record_step("analysis-action", window, method=method)
    window.action_go.trigger()
    app.processEvents()
    if window.model.get_current_outcome_type() == "diagnostic":
        metrics = next(
            w
            for w in app.topLevelWidgets()
            if isinstance(w, DiagnosticMetricsDialog) and w.isVisible()
        )
        for metric in metrics.SELECTABLE_METRICS:
            checkbox = metrics._metric_checkbox(metric)
            if checkbox.isChecked() != (metric == "dor"):
                mouse_click(checkbox, QtCore.Qt.MouseButton.LeftButton)
        mouse_click(metrics.btn_ok, QtCore.Qt.MouseButton.LeftButton)
        app.processEvents()
    dialog = next(
        w
        for w in app.topLevelWidgets()
        if isinstance(w, analysis_setup_dialog.AnalysisSetupDialog) and w.isVisible()
    )
    label = next(
        label for label, key in dialog.available_method_d.items() if key == method
    )
    dialog.method_cbo_box.setCurrentText(label)
    app.processEvents()
    precision_label = next(
        label
        for label in dialog.findChildren(QtWidgets.QLabel)
        if label.text().casefold() == "decimal places"
    )
    precision = precision_label.buddy()
    assert isinstance(precision, QtWidgets.QSpinBox)
    precision.setValue(3)
    record_step("analysis-configured", dialog, method=method, digits=3)
    button = dialog.buttonBox.button(QtWidgets.QDialogButtonBox.StandardButton.Ok)
    assert button.isEnabled()
    mouse_click(required(button), QtCore.Qt.MouseButton.LeftButton)
    app.processEvents()
    views = [
        w
        for w in app.topLevelWidgets()
        if isinstance(w, results_window.ResultsWindow) and w.isVisible()
    ]
    assert len(views) == 1, "analysis did not display results"
    view = views[0]
    assert view.scene.items(), "results scene is empty"
    assert any(
        isinstance(item, QtWidgets.QGraphicsTextItem) and item.toPlainText().strip()
        for item in view.scene.items()
    )
    for path in view.results.display_images.values():
        assert QSvgRenderer(str(path)).isValid(), path
    record_step(
        "analysis-results",
        view,
        texts=dict(view.texts),
        plots=dict(view.results.display_images),
    )
    return view


def analysis_dialog(app, window):
    QtCore.QSettings().remove("analysis/advancedExpanded")
    for sample in ("amino", "continuous", "Pignon2000-HR", "lymph"):
        assert window.open(str(ROOT / f"sample_projects/{sample}.rcms"), raise_on_error=True)
        for row in range(window.model.rowCount()):
            texts = [str(window.model.data(window.model.index(row, column))) for column in window.model.OUTCOMES[:3]]
            decimals = [len(text.split(".")[1]) for text in texts if "." in text and text.replace(".", "", 1).replace("-", "", 1).isdigit()]
            assert len(set(decimals)) <= 1, texts
        kwargs = {"diagnostic_metrics": ["dor"]} if sample == "lymph" else {}
        form = analysis_setup_dialog.AnalysisSetupDialog(
            window.model, parent=window, confidence_level=95, **kwargs)
        window.resize(1600, 828)
        form.resize(631, 437)
        form.show()
        app.processEvents()
        if sample == "amino":
            giamarellou = next(row for row in range(window.model.rowCount()) if "Giamarellou" in str(window.model.data(window.model.index(row, window.model.NAME))))
            assert str(window.model.data(window.model.index(giamarellou, window.model.OUTCOMES[1]))) == "<0.01"
            assert window.model.format_float(0.000401, 3, ratio=True) == "<0.001"
            assert window.tableView.columnWidth(window.model.columnCount() - 1) < 250
            record_step("slice2c-main", window)
            record_step("slice2c-analysis-collapsed", form)
        assert any(label.text() == "Effect measure" and label.isVisible() for label in form.findChildren(QtWidgets.QLabel))
        toggle = form.advanced_toggle
        assert toggle.text().startswith("▸ Advanced")
        assert isinstance(toggle, QtWidgets.QToolButton)
        assert not toggle.isChecked()
        collapsed_height = form.height()
        collapsed_position = form.pos()
        toggle.setFocus()
        key_click(toggle, QtCore.Qt.Key.Key_Space)
        assert toggle.isChecked() and form.advanced_panel.isVisible()
        app.processEvents()
        assert form.height() > collapsed_height, (form.height(), collapsed_height)
        assert form.pos() == collapsed_position
        assert toggle.text().startswith("▾ Advanced")
        record_step("slice2c-analysis-expanded", form)
        form.hide()
        reopened = analysis_setup_dialog.AnalysisSetupDialog(
            window.model, parent=window, confidence_level=95, **kwargs)
        assert reopened.advanced_toggle.isChecked()
        reopened.close()
        form.show()
        digits = next(label.buddy() for label in form.findChildren(QtWidgets.QLabel)
                      if label.text().casefold() == "decimal places")
        digits.setValue(4)

        inference = next(label.buddy() for label in form.findChildren(QtWidgets.QLabel)
                         if "inference" in label.text().casefold())
        inference.setCurrentIndex((inference.currentIndex() + 1) % inference.count())
        expected_inference = inference.currentText()

        key_click(toggle, QtCore.Qt.Key.Key_Return)
        assert not toggle.isChecked()
        assert "4" in toggle.text() and expected_inference in toggle.text()
        record_step("advanced-collapsed", form, sample=sample, summary=toggle.text())

        mouse_click(form.buttonBox.button(QtWidgets.QDialogButtonBox.StandardButton.Ok), QtCore.Qt.MouseButton.LeftButton)
        app.processEvents()
        view = next(w for w in app.topLevelWidgets() if isinstance(w, results_window.ResultsWindow) and w.isVisible())
        assert expected_inference in view.results.details
        record_step("advanced-result", view, configuration=view.results.details)
        view.close()


def sample_analysis(app, window, sample):
    assert window.open(
        str(ROOT / f"sample_projects/{sample}.rcms"), raise_on_error=True
    )
    first_study, study_count = {
        "amino": ("Gonzalez", 19),
        "continuous": ("Carroll", 6),
        "lymph": ("Kinderman", 17),
    }[sample]
    assert (
        str(window.model.data(window.model.index(0, window.model.NAME))) == first_study
    )
    assert sum(study.include for study in window.model.get_studies(False)) == study_count
    method, case = {
        "amino": ("binary.random", "amino-binary-random"),
        "continuous": ("continuous.random", "continuous-random"),
        "lymph": ("diagnostic.random", "lymph-diagnostic-random-dor"),
    }[sample]
    view = analyze(app, window, method)
    contract = json.loads(
        (ROOT / "tests/analysis_regression/baseline/numeric-contract.json").read_text()
    )
    section = "Odds Ratio Summary" if sample == "lymph" else "Summary"
    expected_sections = next(c for c in contract["cases"] if c["id"] == case)[
        "sections"
    ]
    actual_sections = parsed_numeric_sections(view.results)
    expected = expected_sections[section]
    actual = actual_sections[section]
    policy = contract["tolerance_policy"]
    for title, values in expected_sections.items():
        assert title in actual_sections, (title, actual_sections)
        for key, value in values.items():
            tolerance = max(policy["absolute"], policy["relative"] * abs(value))
            assert key in actual_sections[title], (title, key)
            assert abs(actual_sections[title][key] - value) <= tolerance + 1e-12, (
                title,
                key,
                actual_sections[title][key],
                value,
            )
    summary_section = next(item for item in view.results.sections if item.title == section)
    assert summary_section.tables, "summary has no structured table"
    view.report_action.setChecked(True)
    app.processEvents()
    assert view.report_text.isVisible()
    rendered_text = view.report_text.toPlainText()
    for row in summary_section.tables[0].rows:
        assert all(cell in rendered_text for cell in row), (row, rendered_text)
    record_step(
        "numerical-oracle",
        view,
        case=case,
        actual=actual,
        expected=expected,
        tolerance=0.001,
        actual_sections=actual_sections,
        expected_sections=expected_sections,
        tolerance_policy=policy,
    )
    assert view.results.display_images, "analysis produced no plot"
    title = next(iter(view.results.display_images))
    artifact = view.create_plot_artifact(
        title,
        view.results.display_images[title],
        params_path=view.results.image_params_paths[title],
    )
    exported = WORK / "forest.pdf"
    with patch.object(
        QtWidgets.QFileDialog, "getSaveFileName", return_value=(str(exported), "")
    ):
        view.save_image_as(artifact, format="pdf")
    assert exported.read_bytes().startswith(b"%PDF-"), "plot export is not a PDF"
    record_step("pdf-exported", view, file=str(exported))
    view.close()


def sheet_name_dialog(app, name):
    def fill():
        dialog = app.activeModalWidget()
        dialog.findChild(QtWidgets.QLineEdit).setText(name)
        mouse_click(dialog.findChild(QtWidgets.QDialogButtonBox).button(QtWidgets.QDialogButtonBox.StandardButton.Ok), QtCore.Qt.MouseButton.LeftButton)
    QtCore.QTimer.singleShot(0, fill)


def sheets_persist(app, window):
    import copy
    from rc_metastudio.project_adapter import dataset_to_project
    assert window.open(str(ROOT / "sample_projects/amino.rcms"), raise_on_error=True)
    sheet_name_dialog(app, "Third")
    mouse_click(window.add_sheet_button, QtCore.Qt.MouseButton.LeftButton)
    tabs = window.sheet_tabs
    tabs.setCurrentIndex(0)
    table = window.tableView
    model = window.model
    original = copy.deepcopy(dataset_to_project(model.dataset)["sheets"][0])
    table.setFocus()
    selection = QtCore.QItemSelection(model.index(0, model.NAME), model.index(0, model.columnCount() - 1))
    table.selectionModel().select(selection, QtCore.QItemSelectionModel.SelectionFlag.ClearAndSelect)
    key_click(table, QtCore.Qt.Key.Key_C, QtCore.Qt.KeyboardModifier.ControlModifier)
    assert app.clipboard().text()
    tabs.setCurrentIndex(2)
    table.setCurrentIndex(model.index(0, model.NAME))
    table.setFocus()
    key_click(table, QtCore.Qt.Key.Key_V, QtCore.Qt.KeyboardModifier.ControlModifier)
    app.processEvents()
    assert model.rowCount() > 1
    table.setCurrentIndex(model.index(0, model.NAME))
    table.edit(table.currentIndex())
    app.processEvents()
    editor = table.findChild(QtWidgets.QLineEdit)
    editor.setText("Independent pasted row")
    key_click(editor, QtCore.Qt.Key.Key_Return)
    assert dataset_to_project(model.dataset)["sheets"][0] == original
    tabs.setCurrentIndex(0)
    def covariate():
        dialog = app.activeModalWidget()
        dialog.covariate_name_le.setText("Sheet local dose")
        dialog.accept()
    QtCore.QTimer.singleShot(0, covariate)
    window.action_add_covariate.trigger()
    column = table.get_covariate_columns()[-1]
    assert model.setData(model.index(0, column), "7")
    window.action_swap_arms.trigger()
    tabs.setCurrentIndex(1)
    assert "Sheet local dose" not in model.get_covariate_names()
    expected = copy.deepcopy(dataset_to_project(model.dataset))
    current_name = tabs.tabText(1)
    destination = WORK / "persist-three-sheets.rcms"
    with patch.object(QtWidgets.QFileDialog, "getSaveFileName", return_value=(str(destination), "")):
        window.action_save_as.trigger()
    assert window.open(str(destination), raise_on_error=True)
    assert dataset_to_project(model.dataset) == expected
    assert tabs.tabText(tabs.currentIndex()) == current_name
    tabs.setCurrentIndex(2)
    model = window.model
    assert model.data(model.index(0, model.NAME)) == "Independent pasted row"
    record_step("independent-row-and-sheet-state-reopened", window, project=str(destination))


def sheets_edges(app, window):
    sheets(app, window)
    tabs = window.sheet_tabs
    def reject_name(name):
        def fill():
            dialog = app.activeModalWidget()
            dialog.findChild(QtWidgets.QLineEdit).setText(name)
            mouse_click(dialog.findChild(QtWidgets.QDialogButtonBox).button(QtWidgets.QDialogButtonBox.StandardButton.Ok), QtCore.Qt.MouseButton.LeftButton)
            assert dialog.isVisible()
            assert any(label.text() == "Choose a unique outcome name." for label in dialog.findChildren(QtWidgets.QLabel))
            record_step("invalid-sheet-name-rejected", dialog, name=name)
            dialog.reject()
        QtCore.QTimer.singleShot(0, fill)
        window._rename_sheet(0)
    reject_name("")
    reject_name(tabs.tabText(1))
    tabs.setCurrentIndex(0)
    count = tabs.count()
    for answer in (QtWidgets.QMessageBox.StandardButton.No, QtWidgets.QMessageBox.StandardButton.Yes):
        def confirm():
            dialog = app.activeModalWidget()
            assert isinstance(dialog, QtWidgets.QMessageBox)
            dialog.setProperty("verificationExpectedMessage", True)
            record_step("populated-delete-confirmation", dialog, answer=int(answer))
            mouse_click(dialog.button(answer), QtCore.Qt.MouseButton.LeftButton)
        QtCore.QTimer.singleShot(0, confirm)
        window._delete_sheet(0)
        assert tabs.count() == count - (answer == QtWidgets.QMessageBox.StandardButton.Yes)
    tabs.setCurrentIndex(tabs.count() - 1)
    window.action_undo.trigger()
    assert tabs.currentIndex() == 0
    window.action_redo.trigger()
    sheet_name_dialog(app, "Empty deletion")
    mouse_click(window.add_sheet_button, QtCore.Qt.MouseButton.LeftButton)
    with patch.object(QtWidgets.QMessageBox, "question", side_effect=AssertionError("empty delete prompted")):
        window._delete_sheet(tabs.currentIndex())
    while tabs.count() > 1:
        with patch.object(QtWidgets.QMessageBox, "question", return_value=QtWidgets.QMessageBox.StandardButton.Yes):
            window._delete_sheet(0)
    window._delete_sheet(0)
    assert tabs.count() == 1
    record_step("last-sheet-retained", window)


def sheets_swap(app, window):
    import copy
    import math
    from tests.analysis_regression.golden.support.golden_analysis import parsed_numeric_sections
    for filename, metric in (("amino.rcms", "OR"), ("amino.rcms", "RD"), ("continuous.rcms", "MD"), ("Pignon2000-HR.rcms", "HR"), ("NielWeise2008-IRR.rcms", "IRR")):
        window.workspace.mark_saved()
        assert window.open(str(ROOT / "sample_projects" / filename), raise_on_error=True)
        model = window.model
        model.set_current_metric(metric)
        precise = []
        arms = [a.name for a in model.current_sheet.arms]
        favours = copy.deepcopy(model.current_sheet.settings.get("favours_labels"))
        counts = copy.deepcopy([r.analysis_unit.get_raw_data_for_arms(arms) for r in model.current_sheet.rows])
        headers = [model.headerData(i, QtCore.Qt.Orientation.Horizontal) for i in range(model.columnCount())]
        values = []
        for iteration in range(3):
            assert window.action_swap_arms.isEnabled()
            view = analyze(app, window, model.get_current_outcome_type() + ".random")
            if metric in ("OR", "RD", "MD"):
                from rc_metastudio import r_bridge
                r_bridge.load_vars_for_plot(next(iter(view.results.image_params_paths.values())))
                raw = ('metafor::escalc(params$measure, ai=d@g1O1,bi=d@g1O2,ci=d@g2O1,di=d@g2O2,add=params$adjust,to=params$to,drop00=FALSE)' if metric != "MD" else
                       'metafor::escalc("MD",m1i=d@mean1,sd1i=d@sd1,n1i=d@N1,m2i=d@mean2,sd2i=d@sd2,n2i=d@N2)')
                receipt = json.loads(str(r_bridge.execute_r_string('jsonlite::toJSON(local({d<-result$input_data; e<-'+raw+'; ref<-metafor::rma.uni(e$yi,vi=e$vi,method="DL"); actual<-result$res; stopifnot(max(abs(c(actual$b,actual$ci.lb,actual$ci.ub,actual$pval)-c(ref$b,ref$ci.lb,ref$ci.ub,ref$pval)))<1e-6); c(actual$b,actual$ci.lb,actual$ci.ub,actual$pval)}),digits=16)')[0]))
                observed_labels = list(r_bridge.execute_r_string('c(params$fp_favours_left,params$fp_favours_right)'))
                expected_labels = model.current_sheet.settings.get("favours_labels", ["Favours " + name for name in model.arm_names])
                assert observed_labels == expected_labels, (observed_labels, expected_labels)
                precise.append(receipt)
                if iteration == 1:
                    before = precise[0]
                    assert all(math.isclose(x,y,abs_tol=1e-6) for x,y in zip(receipt,[-before[0],-before[2],-before[1],before[3]])), precise
                record_step("swap-direct-metafor", view, metric=metric, iteration=iteration, observed=receipt)
            sections = parsed_numeric_sections(view.results)
            summary = sections["Summary"]
            values.append(summary["estimate"])
            if favours is not None:
                assert model.current_sheet.settings["favours_labels"] == (list(reversed(favours)) if iteration == 1 else favours)
            record_step("swap-pooled-analysis", view, file=filename, iteration=iteration, sections=sections)
            if filename == "amino.rcms" and iteration == 1:
                from xml.etree import ElementTree as ET
                for title, params_path in view.results.image_params_paths.items():
                    exported = WORK / "amino-swapped-forest.svg"
                    view.plot_service.export(regenerator="forest", params_path=params_path, output_path=str(exported))
                    labels = ["".join(node.itertext()) for node in ET.parse(exported).iter("{http://www.w3.org/2000/svg}text")]
                    assert all(name in labels for name in arms), labels
                    record_step("swapped-forest-arm-labels", view, labels=labels, plots={title: str(exported)})
            view.close()
            assert [r.analysis_unit.get_raw_data_for_arms(arms) for r in model.current_sheet.rows] == counts
            if iteration == 0:
                window.action_swap_arms.trigger()
                swapped = [model.headerData(i, QtCore.Qt.Orientation.Horizontal) for i in range(model.columnCount())]
                if filename != "Pignon2000-HR.rcms":
                    assert swapped != headers
                assert [a.name for a in model.current_sheet.arms] == list(reversed(arms))
            elif iteration == 1:
                window.action_swap_arms.trigger()
                assert [model.headerData(i, QtCore.Qt.Orientation.Horizontal) for i in range(model.columnCount())] == headers
        assert values[2] == values[0], values
        if metric in ("RD", "MD"):
            assert abs(values[0] + values[1]) <= .002, values
        else:
            assert math.isclose(values[1], 1 / values[0], abs_tol=.003), values
    window.workspace.mark_saved()
    assert window.open(str(ROOT / "sample_projects/lymph.rcms"), raise_on_error=True)
    assert not window.action_swap_arms.isEnabled()
    sheet_name_dialog(app, "Diagnostic 2")
    mouse_click(window.add_sheet_button, QtCore.Qt.MouseButton.LeftButton)
    for index in (0, 1, 0):
        window.sheet_tabs.setCurrentIndex(index)
        assert not window.action_swap_arms.isEnabled()
    record_step("diagnostic-swap-disabled", window)


def sheets_legacy(app, window):
    import hashlib
    import subprocess
    from tests.analysis_regression.golden.support.golden_analysis import parsed_numeric_sections
    baseline = json.loads((ROOT / "tests/e2e/fixtures/sheets-original-numerics.json").read_text())
    for filename in dict.fromkeys(row["file"] for row in baseline):
        window.workspace.mark_saved()
        source = WORK / filename
        source.write_bytes(subprocess.check_output(["git", "show", "322ede26:sample_projects/" + filename], cwd=ROOT))
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        assert window.open(str(source), raise_on_error=True)
        rows = [row for row in baseline if row["file"] == filename]
        names = [row["pair"]["outcome"] if sum(r["pair"]["outcome"] == row["pair"]["outcome"] for r in rows) == 1 else row["pair"]["outcome"] + " — " + row["pair"]["follow_up"] for row in rows]
        assert [window.sheet_tabs.tabText(i) for i in range(window.sheet_tabs.count())] == names
        for index, row in enumerate(rows):
            window.sheet_tabs.setCurrentIndex(index)
            view = analyze(app, window, row["method"])
            actual = parsed_numeric_sections(view.results)
            assert actual == row["sections"], (filename, actual, row["sections"])
            record_step("legacy-original-numerics", view, file=filename, sha256=digest, sections=actual)
            view.close()
        assert hashlib.sha256(source.read_bytes()).hexdigest() == digest


def sheets(app, window):
    import copy
    from rc_metastudio.project_adapter import dataset_to_project
    def project(dataset):
        return dataset_to_project(dataset)
    assert window.open(str(ROOT / "sample_projects/amino.rcms"), raise_on_error=True)
    tabs = window.sheet_tabs
    original = copy.deepcopy(window.model.dataset)
    window.resize(1600, 828)
    app.processEvents()
    record_step("sheets-initial", window)
    screens = WORK / "sheets-screens"
    screens.mkdir(exist_ok=True)
    window.grab().save(str(screens / "sheets-after-main.png"))
    window.action_go.trigger()
    app.processEvents()
    dialog = next(w for w in app.topLevelWidgets() if isinstance(w, analysis_setup_dialog.AnalysisSetupDialog) and w.isVisible())
    dialog.resize(631, 437)
    app.processEvents()
    dialog.grab().save(str(screens / "sheets-after-analysis.png"))
    dialog.close()

    def name_dialog(name, duplicate=False):
        def fill():
            dialog = app.activeModalWidget()
            edit = dialog.findChild(QtWidgets.QLineEdit)
            edit.setText(name)
            ok = dialog.findChild(QtWidgets.QDialogButtonBox).button(QtWidgets.QDialogButtonBox.StandardButton.Ok)
            mouse_click(ok, QtCore.Qt.MouseButton.LeftButton)
            if duplicate:
                assert dialog.isVisible()
                assert any(label.text() == "Choose a unique outcome name." for label in dialog.findChildren(QtWidgets.QLabel))
                record_step("sheets-duplicate-rejected", dialog)
                dialog.reject()
        QtCore.QTimer.singleShot(0, fill)

    for name in ("Outcome 2", "Outcome 3"):
        name_dialog(name)
        mouse_click(window.add_sheet_button, QtCore.Qt.MouseButton.LeftButton)
        assert window.model.current_sheet.name == name
        assert not window.model.current_sheet.rows
        tabs.setCurrentIndex(0)
        window.action_undo.trigger()
        assert name not in [s.name for s in window.model.dataset.sheets]
        window.action_redo.trigger()
    name_dialog("Renamed outcome")
    QTest.mouseDClick(tabs, QtCore.Qt.MouseButton.LeftButton, pos=tabs.tabRect(3).center())
    assert window.model.current_sheet.name == "Renamed outcome"
    tabs.setCurrentIndex(0)
    window.action_undo.trigger()
    assert window.model.current_sheet.name == "Outcome 3"
    window.action_redo.trigger()
    name_dialog(original.sheets[0].name, duplicate=True)
    QTest.mouseDClick(tabs, QtCore.Qt.MouseButton.LeftButton, pos=tabs.tabRect(3).center())
    tabs.setCurrentIndex(2)
    table = window.tableView
    table.setCurrentIndex(window.model.index(0, window.model.NAME))
    table.edit(table.currentIndex())
    app.processEvents()
    editor = table.findChild(QtWidgets.QLineEdit)
    editor.setText("Independent row")
    key_click(editor, QtCore.Qt.Key.Key_Return)
    assert len(window.model.current_sheet.rows) == 1
    tabs.setCurrentIndex(0)
    assert project(window.model.dataset)["sheets"][0] == project(original)["sheets"][0]
    window.action_undo.trigger()
    assert window.model.current_sheet.name == "Outcome 2"
    assert not window.model.current_sheet.rows
    window.action_redo.trigger()
    tabs.setCurrentIndex(0)
    window.activateWindow()
    table.setFocus()
    app.processEvents()
    key_click(table, QtCore.Qt.Key.Key_PageDown, QtCore.Qt.KeyboardModifier.ControlModifier)
    assert tabs.currentIndex() == 1
    QTest.mousePress(tabs, QtCore.Qt.MouseButton.LeftButton, pos=tabs.tabRect(2).center())
    QTest.mouseMove(tabs, tabs.tabRect(3).center(), 100)
    QTest.mouseRelease(tabs, QtCore.Qt.MouseButton.LeftButton, pos=tabs.tabRect(3).center())
    assert tabs.tabText(3) == "Outcome 2"
    tabs.setCurrentIndex(0)
    window.action_undo.trigger()
    assert window.model.current_sheet.name == "Outcome 2", (window.model.current_sheet.name, [s.name for s in window.model.dataset.sheets])
    window.action_redo.trigger()
    record_step("sheets-reordered", window)
    index = tabs.currentIndex()
    def delete_menu():
        menu = app.activePopupWidget()
        action = next(a for a in menu.actions() if a.text() == "Delete")
        QtCore.QTimer.singleShot(0, lambda: app.activeModalWidget().done(QtWidgets.QMessageBox.StandardButton.Yes))
        mouse_click(menu, QtCore.Qt.MouseButton.LeftButton, menu.actionGeometry(action).center())
    QtCore.QTimer.singleShot(0, delete_menu)
    window._sheet_menu(tabs.tabRect(index).center())
    tabs.setCurrentIndex(0)
    window.action_undo.trigger()
    assert window.model.current_sheet.name == "Outcome 2", (window.model.current_sheet.name, [s.name for s in window.model.dataset.sheets])
    assert len(window.model.current_sheet.rows) == 1
    tabs.setCurrentIndex(0)
    before = copy.deepcopy(window.model.current_sheet)
    effect = window.model.get_current_arm_key()
    metric = window.model.current_effect
    unit = window.model.current_sheet.rows[0].analysis_unit
    value = unit.get_display_effect_for_source("derived_preview", metric, effect)
    counts = copy.deepcopy(unit.get_raw_data_for_arms([a.name for a in before.arms]))
    window.action_swap_arms.trigger()
    assert window.model.current_sheet.arms == list(reversed(before.arms))
    assert [r.label for r in window.model.current_sheet.rows] == [r.label for r in before.rows]
    assert window.model.get_current_arm_key() != effect
    unit = window.model.current_sheet.rows[0].analysis_unit
    reversed_value = unit.get_display_effect_for_source("derived_preview", metric, window.model.get_current_arm_key())
    assert abs(value * reversed_value - 1) < 1e-8, (value, reversed_value)
    assert unit.get_raw_data_for_arms([a.name for a in before.arms]) == counts
    record_step("sheets-swapped", window, original_effect=value, reversed_effect=reversed_value, counts=counts)
    tabs.setCurrentIndex(1)
    window.action_undo.trigger()
    assert project(window.model.dataset)["sheets"][0] == project(original)["sheets"][0]
    window.action_redo.trigger()
    window.action_swap_arms.trigger()
    actual_sheet = project(window.model.dataset)["sheets"][0]
    expected_sheet = project(original)["sheets"][0]
    assert actual_sheet == expected_sheet, {k: (expected_sheet[k], v) for k, v in actual_sheet.items() if v != expected_sheet[k]}
    window.action_swap_arms.trigger()
    expected = copy.deepcopy(window.model.dataset)
    destination = WORK / "sheets.rcms"
    with patch.object(QtWidgets.QFileDialog, "getSaveFileName", return_value=(str(destination), "")):
        window.action_save_as.trigger()
    assert window.open(str(destination), raise_on_error=True)
    assert project(window.model.dataset) == project(expected)
    record_step("sheets-reopened", window, names=[s.name for s in expected.sheets])


def editing(app, window):
    assert window.open(str(ROOT / "sample_projects/amino.rcms"), raise_on_error=True)
    table = window.tableView
    original = str(window.model.data(window.model.index(0, window.model.NAME)))
    record_step("table-before-edit", window, name=original)
    table.setCurrentIndex(window.model.index(0, window.model.NAME))
    table.edit(table.currentIndex())
    app.processEvents()
    editor = table.findChild(QtWidgets.QLineEdit)
    assert editor is not None
    editor.setText("München – edited")
    key_click(editor, QtCore.Qt.Key.Key_Return)
    app.processEvents()
    assert (
        str(window.model.data(window.model.index(0, window.model.NAME)))
        == "München – edited"
    )
    assert table.currentIndex().row() == 1
    assert window.workspace.is_dirty
    window.action_undo.trigger()
    assert str(window.model.data(window.model.index(0, window.model.NAME))) == original
    assert not window.workspace.is_dirty
    window.action_redo.trigger()
    destination = WORK / "edited.rcms"
    with patch.object(
        QtWidgets.QFileDialog, "getSaveFileName", return_value=(str(destination), "")
    ):
        assert window.save_as()
    assert not window.workspace.is_dirty
    assert window.open(str(destination), raise_on_error=True)
    assert (
        str(window.model.data(window.model.index(0, window.model.NAME)))
        == "München – edited"
    )
    # A malformed project must leave the current document and save path intact.
    broken = WORK / "broken.rcms"
    broken.write_bytes(b"not a project")
    with patch.object(
        QtWidgets.QMessageBox,
        "critical",
        return_value=QtWidgets.QMessageBox.StandardButton.Ok,
    ) as error:
        assert not window.open(str(broken))
    assert error.called
    assert (
        str(window.model.data(window.model.index(0, window.model.NAME)))
        == "München – edited"
    )
    assert window.out_path == str(destination)
    record_step(
        "edit-reopened",
        window,
        name="München – edited",
        project=str(destination),
        corrupt_open_preserved=True,
    )
    analyze(app, window, "binary.random").close()


def xlsx_import(app, window):
    from rc_metastudio import main_wizard
    original_palette = app.palette()
    screens = ROOT / "docs/ux-hill-climb/screens"
    screens.mkdir(parents=True, exist_ok=True)
    for theme, foreground, background in (("light", "#202020", "#ffffff"), ("dark", "#eeeeee", "#202020")):
        palette = QtGui.QPalette(original_palette)
        for role in (QtGui.QPalette.ColorRole.Text, QtGui.QPalette.ColorRole.ButtonText, QtGui.QPalette.ColorRole.WindowText):
            palette.setColor(role, QtGui.QColor(foreground))
        for role in (QtGui.QPalette.ColorRole.Window, QtGui.QPalette.ColorRole.Button, QtGui.QPalette.ColorRole.Base):
            palette.setColor(role, QtGui.QColor(background))
        app.setPalette(palette)
        capture = main_wizard.MainWizard(path="new_dataset")
        capture.show()
        app.processEvents()
        types = capture.page(main_wizard.Page_DataType)
        buttons = types._data_type_buttons()
        assert all(not button.icon().isNull() for button in buttons)
        capture.grab().save(str(screens / f"slice7-{theme}-data-types.png"))
        capture.grab().save(str(screens / f"slice8a-data-types-{theme}.png"))
        for page_id, page_name in ((main_wizard.Page_Welcome, "welcome"),
                                   (main_wizard.Page_OutcomeName, "outcome"),
                                   (main_wizard.Page_CsvImport, "csv")):
            capture.setStartId(page_id)
            capture.restart()
            app.processEvents()
            capture.grab().save(str(screens / f"slice8a-{page_name}-{theme}.png"))
        window.grab().save(str(screens / f"slice8a-grid-{theme}.png"))
        for index, button in enumerate(buttons):
            types._button_selected(button)
            if capture.require_dataset_info()["data_type"] == "diagnostic":
                continue
            capture.setStartId(main_wizard.Page_ChooseMetric)
            capture.restart()
            app.processEvents()
            metrics = capture.currentPage().metric_cbo_box
            assert all(metrics.itemIcon(i).isNull() for i in range(metrics.count()))
            capture.grab().save(str(screens / f"slice7-{theme}-metrics-{index}.png"))
            if theme == "light" and index == 0:
                capture.grab().save(str(screens / "slice8a-metric.png"))
        capture.close()
    app.setPalette(original_palette)
    from openpyxl import Workbook
    from rc_metastudio import main_wizard
    book = Workbook()
    book.active.title = "Empty"
    sheet = book.create_sheet("Binary")
    sheet.append(["Study Name", "Year", "Tx A #evts", "Tx A #total", "Tx B #evts", "Tx B #total", "OR", "Lower", "Upper"])
    rows = [["Alpha", 2020, 1, 10, 2, 12, None, None, None], ["Beta", 2021, 3, 11, 4, 13, None, None, None]]
    for row in rows:
        sheet.append(row)
    hidden = book.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["D"].hidden = True
    path = WORK / "studies.xlsx"
    book.save(path)
    from rc_metastudio.project_adapter import dataset_to_project
    unchanged = dataset_to_project(window.model.dataset)
    with patch.object(QtWidgets.QFileDialog, "getOpenFileName", return_value=("", "")):
        window.action_import_csv.trigger()
    assert dataset_to_project(window.model.dataset) == unchanged
    def inspect_import(dialog):
        loaded = dialog.currentPage()
        assert isinstance(loaded, main_wizard.CsvImportPage)
        assert loaded.file_path == str(path)
        assert loaded.isComplete(), loaded.import_status.text()
        record_step("toolbar-import-loaded", dialog)
        return 0
    with patch.object(QtWidgets.QFileDialog, "getOpenFileName", return_value=(str(path), "")), patch.object(main_wizard.MainWizard, "exec", inspect_import):
        window.action_import_csv.trigger()
    assert dataset_to_project(window.model.dataset) == unchanged
    welcome = main_wizard.MainWizard()
    welcome.show()
    app.processEvents()
    with patch.object(QtWidgets.QFileDialog, "getOpenFileName", return_value=(str(path), "")):
        mouse_click(welcome.page(main_wizard.Page_Welcome).import_csv_btn, QtCore.Qt.MouseButton.LeftButton)
    assert isinstance(welcome.currentPage(), main_wizard.CsvImportPage)
    assert welcome.currentPage().file_path == str(path)
    assert welcome.currentPage().isComplete(), welcome.currentPage().import_status.text()
    record_step("welcome-import-loaded", welcome)
    welcome.grab().save(str(screens / "slice8a-import.png"))
    welcome.close()
    wizard = main_wizard.MainWizard(path="csv_import")
    wizard.set_dataset_info(dict(arms="two", data_type="binary", sub_type="proportions", effect="OR", metric_choices=[], name="Mortality"))
    wizard.page(main_wizard.Page_OutcomeName).outcome_name_LineEdit.setText("Mortality")
    wizard.setStartId(main_wizard.Page_CsvImport)
    wizard.restart()
    wizard.show()
    app.processEvents()
    page = wizard.currentPage()
    with patch.object(QtWidgets.QFileDialog, "getOpenFileName", return_value=(str(path), "")):
        mouse_click(page.select_file_btn, QtCore.Qt.MouseButton.LeftButton)
    assert page.isComplete(), page.import_status.text()
    assert page.worksheet_combo.count() == 2
    assert page.worksheet_combo.currentText() == "Binary"
    assert "1 hidden row was imported" in page.import_status.text()
    assert "1 hidden column was imported" in page.import_status.text()
    assert page.title() == "Import data"
    assert not page.csv_options_btn.isVisible()
    assert page.preview_table.isVisible()
    assert not page.pageScrollArea.verticalScrollBar().isVisible()
    payload = page.csv_data()
    import csv as csv_module
    csv_path = WORK / "equivalent.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv_module.writer(stream)
        writer.writerow([c.value for c in sheet[1]])
        writer.writerows(rows)
    from rc_metastudio.csv_import import parse_csv
    equivalent = parse_csv(csv_path, expected_headers=page.required_header_labels, has_headers=True, from_excel=True, header_aliases=page._legacy_header_aliases())
    assert payload == equivalent.to_payload()
    from rc_metastudio.project_adapter import dataset_to_project
    before = dataset_to_project(window.model.dataset)
    def select(candidate):
        with patch.object(QtWidgets.QFileDialog, "getOpenFileName", return_value=(str(candidate), "")):
            mouse_click(page.select_file_btn, QtCore.Qt.MouseButton.LeftButton)
    for filename, value, expected in [("formula.xlsx", "=1+1", "Binary!C2"), ("error.xlsx", "#DIV/0!", "Binary!C2")]:
        sheet["C2"] = value
        candidate = WORK / filename
        book.save(candidate)
        select(candidate)
        assert not page.isComplete()
        assert expected in page.import_status.text(), page.import_status.text()
        assert dataset_to_project(window.model.dataset) == before
        record_step(filename, wizard, error=page.import_status.text())
    sheet["C2"] = 1
    sheet.merge_cells("C2:D2")
    candidate = WORK / "merged.xlsx"
    book.save(candidate)
    select(candidate)
    assert not page.isComplete()
    assert "C2:D2" in page.import_status.text()
    for filename in ("old.xls", "corrupt.xlsx", "encrypted.xlsx"):
        candidate = WORK / filename
        candidate.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") if filename == "encrypted.xlsx" else b"not a workbook")
        select(candidate)
        assert not page.isComplete()
        assert dataset_to_project(window.model.dataset) == before
    # Excel stores cached formula values in the worksheet XML; openpyxl does
    # not calculate formulas, so write the cache explicitly in this fixture.
    from zipfile import ZipFile, ZIP_DEFLATED
    formula_path = WORK / "formula.xlsx"
    cached_path = WORK / "cached.xlsx"
    with ZipFile(formula_path) as source, ZipFile(cached_path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet2.xml":
                content = content.replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
            target.writestr(item, content)
    select(cached_path)
    assert page.isComplete(), page.import_status.text()
    assert "1 formula cells use their saved values" in page.import_status.text()
    assert page.csv_data()["data"][0][2] == "2"
    select(path)

    assert payload["data"][0][:6] == ["Alpha", "2020", "1", "10", "2", "12"]
    record_step("xlsx-preview", wizard, file=str(path), data=payload)
    page.worksheet_combo.setCurrentText("Empty")
    assert not page.isComplete()
    page.worksheet_combo.setCurrentText("Binary")
    assert page.isComplete()
    window._handle_wizard_results(wizard.get_results())
    imported_project = dataset_to_project(window.model.dataset)
    window.action_undo.trigger()
    assert dataset_to_project(window.model.dataset) == before
    window.action_redo.trigger()
    assert dataset_to_project(window.model.dataset) == imported_project
    wizard.close()
    view = analyze(app, window, "binary.random")
    xlsx_results = dict(view.texts)
    view.close()
    window.out_path = str(WORK / "xlsx.rcms")
    assert window.save()
    assert window.open(window.out_path, raise_on_error=True)
    assert window.model.get_studies(False)[0].label == "Alpha"
    record_step("xlsx-reopened", window, project=window.out_path)
    select(csv_path)
    assert page.isComplete(), page.import_status.text()
    window._handle_wizard_results(wizard.get_results())
    def normalize_ids(value):
        import re
        text = json.dumps(value)
        identifiers = re.findall(r'[0-9a-f]{32}', text)
        for index, identifier in enumerate(dict.fromkeys(identifiers)):
            text = text.replace(identifier, f"id-{index}")
        return json.loads(text)
    (WORK / "csv-project.json").write_text(json.dumps(dataset_to_project(window.model.dataset), indent=2))
    (WORK / "xlsx-project.json").write_text(json.dumps(imported_project, indent=2))
    assert normalize_ids(dataset_to_project(window.model.dataset)) == normalize_ids(imported_project)
    view = analyze(app, window, "binary.random")
    assert dict(view.texts) == xlsx_results
    view.close()


    from datetime import datetime
    continuous = Workbook()
    sheet = continuous.active
    sheet.title = "Continuous"
    info = dict(arms="two", data_type="continuous", sub_type="means_and_std_devs", effect="MD", metric_choices=[], name="Precision")
    precision_wizard = main_wizard.MainWizard(path="csv_import")
    precision_wizard.set_dataset_info(info)
    precision_wizard.setStartId(main_wizard.Page_CsvImport)
    precision_wizard.restart()
    precision_wizard.show()
    app.processEvents()
    precision_page = precision_wizard.currentPage()
    sheet.append([])
    sheet.append([None] + precision_page.required_header_labels)
    values = ["Precision", datetime(2022, 1, 1), 12.0, 0.1, 1/3, 12.0, 0.1, 1/3]
    sheet.append([None] + values)
    sheet.row_dimensions[3].hidden = True
    precision_path = WORK / "precision.xlsx"
    continuous.save(precision_path)
    with patch.object(QtWidgets.QFileDialog, "getOpenFileName", return_value=(str(precision_path), "")):
        mouse_click(precision_page.select_file_btn, QtCore.Qt.MouseButton.LeftButton)
    assert precision_page.isComplete(), precision_page.import_status.text()
    assert precision_page.csv_data()["data"][0][1:5] == ["2022", "12", "0.1", repr(1/3)]
    assert "Continuous!C3" in precision_page.import_status.text()
    record_step("xlsx-precision", precision_wizard, data=precision_page.csv_data())
    precision_wizard.close()


def csv_import(app, window):
    from rc_metastudio import main_wizard

    csv = WORK / "studies.csv"
    csv.write_text(
        'Study Name,Year,Tx A #evts,Tx A #total,Tx B #evts,Tx B #total,OR,Lower,Upper,Dose,Region\n"Alpha, München",2020,1,10,2,12,,,,5.5,North\nBeta,2021,3,11,4,13,,,,7,South\nGamma,2022,2,20,5,20,,,,8,West\n',
        encoding="utf-8-sig",
    )
    wizard = main_wizard.MainWizard(path="csv_import")
    wizard.set_dataset_info(
        {
            "arms": "two",
            "data_type": "binary",
            "sub_type": "proportions",
            "effect": "OR",
            "metric_choices": [],
            "name": "Mortality",
        }
    )
    name_page = wizard.page(main_wizard.Page_OutcomeName)
    assert isinstance(name_page, main_wizard.OutcomeNamePage)
    name_page.outcome_name_LineEdit.setText("Mortality")
    wizard.setStartId(main_wizard.Page_CsvImport)
    wizard.restart()
    wizard.show()
    app.processEvents()
    page = wizard.currentPage()
    assert isinstance(page, main_wizard.CsvImportPage)
    with patch.object(
        QtWidgets.QFileDialog, "getOpenFileName", return_value=(str(csv), "")
    ):
        mouse_click(page.select_file_btn, QtCore.Qt.MouseButton.LeftButton)
    assert page.isComplete()
    assert page.preview_table.rowCount() == 3
    assert page.csv_options_btn.isVisible()
    assert not page.groupBox.isVisible()
    mouse_click(page.csv_options_btn, QtCore.Qt.MouseButton.LeftButton)
    assert page.delimter_le.isVisible()
    assert page.has_headers_chkbx.isVisible()
    record_step("csv-preview", wizard, rows=3, file=str(csv))
    window._handle_wizard_results(wizard.get_results())
    wizard.close()
    assert (
        str(window.model.data(window.model.index(0, window.model.NAME)))
        == "Alpha, München"
    )
    assert window.model.get_studies(False)[1].covariate_values[next(c["id"] for c in window.model.current_sheet.covariates if c["name"] == "Dose")] == 7.0
    assert window.model.get_studies(False)[1].covariate_values[next(c["id"] for c in window.model.current_sheet.covariates if c["name"] == "Region")] == "South"
    assert all(study.include for study in window.model.get_studies(False))
    view = analyze(app, window, "binary.random")
    assert "Alpha" in view.texts.get("Weights", "")
    view.close()
    window.out_path = str(WORK / "imported.rcms")
    assert window.save()
    assert window.open(window.out_path, raise_on_error=True)
    assert window.model.get_studies(False)[0].label == "Alpha, München"
    reopened = window.model.get_studies(False)[1]
    assert reopened.covariate_values[next(c["id"] for c in window.model.current_sheet.covariates if c["name"] == "Dose")] == 7.0
    assert reopened.covariate_values[next(c["id"] for c in window.model.current_sheet.covariates if c["name"] == "Region")] == "South"
    assert window.model.current_effect == "OR"
    record_step(
        "csv-reopened",
        window,
        project=window.out_path,
        study="Alpha, München",
        dose=reopened.covariate_values[next(c["id"] for c in window.model.current_sheet.covariates if c["name"] == "Dose")],
        region=reopened.covariate_values[next(c["id"] for c in window.model.current_sheet.covariates if c["name"] == "Region")],
    )


def plot_edit(app, window):
    from rc_metastudio.plot_editor_dialog import EditPlotDialog

    assert window.open(str(ROOT / "sample_projects/amino.rcms"), raise_on_error=True)
    view = analyze(app, window, "binary.random")
    title = next(iter(view.results.display_images))
    path = Path(view.results.display_images[title])
    before = path.read_bytes()
    artifact = view.create_plot_artifact(
        title,
        view.results.images[title],
        params_path=view.results.image_params_paths[title],
    )
    callback_errors = []

    def edit():
        dialog = app.activeModalWidget()
        if not isinstance(dialog, EditPlotDialog):
            QtCore.QTimer.singleShot(10, edit)
            return
        try:
            dialog.x_lbl_le.setText("Edited effect axis")
            record_step("plot-edit-draft", dialog, axis="Edited effect axis")
            mouse_click(
                required(
                    dialog.buttonBox.button(
                        QtWidgets.QDialogButtonBox.StandardButton.Apply
                    )
                ),
                QtCore.Qt.MouseButton.LeftButton,
            )
            assert path.read_bytes() != before
            assert (
                view.plot_service.load_params(artifact.params_path)["fp_xlabel"]
                == "Edited effect axis"
            )
            dialog.x_lbl_le.setText("Discard this draft")
        except Exception as error:  # noqa: BLE001 - propagate Qt callback failures after the modal loop.
            callback_errors.append(error)
        finally:
            dialog.reject()

    QtCore.QTimer.singleShot(0, edit)
    view.edit_plot(artifact, view._svg_plot_items[0])
    assert not callback_errors, callback_errors
    assert (
        view.plot_service.load_params(artifact.params_path)["fp_xlabel"]
        == "Edited effect axis"
    )
    assert QSvgRenderer(str(path)).isValid()
    record_step(
        "plot-applied-cancelled",
        view,
        axis="Edited effect axis",
        plots={title: str(path)},
    )
    view.close()


def small_study_effects(app, window):
    from rc_metastudio.publication_bias_dialog import PublicationBiasDialog

    assert window.open(str(ROOT / "sample_projects/amino.rcms"), raise_on_error=True)
    callback_errors = []

    def run():
        dialog = app.activeModalWidget()
        if not isinstance(dialog, PublicationBiasDialog):
            QtCore.QTimer.singleShot(10, run)
            return
        try:
            button = dialog.button_box.button(
                QtWidgets.QDialogButtonBox.StandardButton.Ok
            )
            assert button.isEnabled()
            if not dialog.ordinary_funnel_check.isChecked():
                dialog.ordinary_funnel_check.setChecked(True)
            record_step("small-study-configured", dialog, ordinary_funnel=True)
            mouse_click(required(button), QtCore.Qt.MouseButton.LeftButton)
            assert not dialog.failure_label.isVisible(), dialog.failure_label.text()
        except Exception as error:  # noqa: BLE001 - propagate Qt callback failures after the modal loop.
            callback_errors.append(error)
        finally:
            dialog.reject()

    QtCore.QTimer.singleShot(0, run)
    window.action_publication_bias.trigger()
    assert not callback_errors, callback_errors
    view = next(
        w
        for w in app.topLevelWidgets()
        if isinstance(w, results_window.ResultsWindow) and w.isVisible()
    )
    assert view.results.images, dict(view.texts)
    assert any("Harbord" in text or "Rücker" in text for text in view.texts.values())
    for path in view.results.images.values():
        if Path(path).suffix == ".svg":
            assert QSvgRenderer(str(path)).isValid()
        else:
            assert not QtGui.QImage(str(path)).isNull()
    record_step(
        "small-study-results",
        view,
        texts=dict(view.texts),
        plots=dict(view.results.images),
    )
    view.close()


def calculator(app, window, family):
    from rc_metastudio import (
        binary_data_dialog,
        continuous_data_dialog,
        diagnostic_data_dialog,
    )

    sample, dialog_type, table_name, column = {
        "binary": ("amino", binary_data_dialog.BinaryDataDialog, "raw_data_table", 0),
        "continuous": (
            "continuous",
            continuous_data_dialog.ContinuousDataDialog,
            "simple_table",
            1,
        ),
        "diagnostic": (
            "lymph",
            diagnostic_data_dialog.DiagnosticDataDialog,
            "two_by_two_table",
            0,
        ),
    }[family]
    assert window.open(
        str(ROOT / f"sample_projects/{sample}.rcms"), raise_on_error=True
    )

    def raw():
        model = window.model
        return list(
            model.get_current_analysis_unit_for_study(0).get_raw_data_for_arm(
                model.arm_names[0]
            )
        )

    before = raw()
    record_step("calculator-before", window, family=family, raw=before)
    expected = before.copy()
    expected[column] += 1.5 if family == "continuous" else 1
    if family == "binary":
        expected[1] += 1
    callback_errors = []

    def edit(accept):
        dialog = app.activeModalWidget()
        if not isinstance(dialog, dialog_type):
            QtCore.QTimer.singleShot(10, lambda: edit(accept))
            return
        try:
            table = getattr(dialog, table_name)
            app.processEvents()
            dialog.grab().save(str(ROOT / f"docs/ux-hill-climb/screens/slice8a-calculator-{family}.png"))
            for field in (dialog.effect_text_box, dialog.lower_text_box, dialog.upper_text_box):
                width = field.fontMetrics().horizontalAdvance(field.text())
                assert field.width() >= width + field.textMargins().left() + field.textMargins().right() + 12
            table.setCurrentCell(0, column)
            table.item(0, column).setText(
                str(expected[column] if accept else expected[column] + 1).replace(
                    ".", ","
                )
            )
            if accept:
                dialog.effect_text_box.setFocus()
                app.processEvents()
                from rc_metastudio.meta_globals import CALC_NUM_DIGITS
                for field in (dialog.effect_text_box, dialog.lower_text_box, dialog.upper_text_box):
                    assert len(field.text().partition(".")[2]) <= CALC_NUM_DIGITS, field.text()
                if family == "continuous":
                    assert not table.horizontalScrollBar().isVisible()
                assert dialog.findChild(QtWidgets.QLabel, "calculator_wald") is not None
                if family == "binary":
                    import math
                    from rpy2 import robjects
                    counts = dialog.analysis_unit.get_raw_data_for_arms(dialog.arm_names)
                    e1, n1, e2, n2 = counts
                    crosscheck = {}
                    for row, metric in enumerate(("OR", "RR", "RD")):
                        result = robjects.r(f'local({{ x <- metafor::escalc(measure="{metric}", ai={e1}, bi={n1-e1}, ci={e2}, di={n2-e2}); se <- sqrt(x$vi); c(x$yi, se, x$yi/se, 2*pnorm(-abs(x$yi/se))) }})')
                        est, se, z, p = map(float, result)
                        transform = math.exp if metric != "RD" else float
                        for col, value in enumerate((transform(est), transform(est-dialog.confidence_multiplier*se), transform(est+dialog.confidence_multiplier*se)), 1):
                            assert dialog.metric_readout.item(row, col).text() == f"{value:.4f}", (metric, col, dialog.metric_readout.item(row, col).text(), value)
                        crosscheck[metric] = dict(estimate=est, se=se, z=z, p=p)
                    for row, metric in enumerate(("OR", "RR", "RD")):
                        table_item = dialog.metric_readout.item(row, 0)
                        mouse_click(dialog.metric_readout.viewport(), QtCore.Qt.MouseButton.LeftButton,
                                    pos=dialog.metric_readout.visualItemRect(table_item).center())
                        assert dialog.current_effect == metric
                        result = crosscheck[metric]
                        text = dialog.wald_readout.text()
                        assert f"{result['se']:.4f}" in text
                        assert f"{result['z']:.4f}" in text
                        assert ("<0.001" if result["p"] < 0.001 else f"{result['p']:.3f}") in text
                    mouse_click(dialog.metric_readout.viewport(), QtCore.Qt.MouseButton.LeftButton,
                                pos=dialog.metric_readout.visualItemRect(dialog.metric_readout.item(0, 0)).center())
                    selected = crosscheck[dialog.current_effect]
                    text = dialog.wald_readout.text()
                    assert f"{selected['se']:.4f}" in text
                    assert f"{selected['z']:.4f}" in text
                    assert ("<0.001" if selected["p"] < 0.001 else f"{selected['p']:.3f}") in text
                    record_step("calculator-r-crosscheck", dialog, counts=counts, results=crosscheck)

                record_step("calculator-effect-preview", dialog, family=family,
                            previews=[field.text() for field in
                                      (dialog.effect_text_box, dialog.lower_text_box, dialog.upper_text_box)])
                dialog.lower_text_box.setFocus()
                app.processEvents()
                record_step("calculator-ci-preview", dialog, family=family,
                            lower=dialog.lower_text_box.text(), upper=dialog.upper_text_box.text())
            button = dialog.buttonBox.button(
                QtWidgets.QDialogButtonBox.StandardButton.Ok
                if accept
                else QtWidgets.QDialogButtonBox.StandardButton.Cancel
            )
            assert button.isEnabled()
            record_step(
                "calculator-accept" if accept else "calculator-cancel",
                dialog,
                family=family,
            )
            mouse_click(required(button), QtCore.Qt.MouseButton.LeftButton)
        except Exception as error:  # noqa: BLE001 - propagate Qt callback failures after the modal loop.
            callback_errors.append(error)
            dialog.reject()

    QtCore.QTimer.singleShot(0, lambda: edit(True))
    header = window.tableView.verticalHeader()
    first_row = QtCore.QPoint(
        header.viewport().width() // 2,
        header.sectionViewportPosition(0) + header.sectionSize(0) // 2,
    )
    mouse_click(header.viewport(), QtCore.Qt.MouseButton.LeftButton, pos=first_row)
    assert not callback_errors, callback_errors
    assert raw() == expected, (raw(), expected)
    assert window.workspace.is_dirty
    QtCore.QTimer.singleShot(0, lambda: edit(False))
    mouse_click(header.viewport(), QtCore.Qt.MouseButton.LeftButton, pos=first_row)
    assert not callback_errors, callback_errors
    assert raw() == expected
    window.action_undo.trigger()
    assert raw() == before
    assert not window.workspace.is_dirty
    window.action_redo.trigger()
    assert raw() == expected
    window.out_path = str(WORK / "calculator.rcms")
    assert window.save()
    import zipfile
    saved_row = json.loads(zipfile.ZipFile(window.out_path).read("project.json"))["sheets"][0]["rows"][0]
    baseline = json.loads((ROOT / "tests/e2e/calculator-accepted-baseline.json").read_text())["rows"][family]
    def float_bits(value):
        if isinstance(value, float):
            return value.hex()
        if isinstance(value, dict):
            return {key: float_bits(val) for key, val in value.items()}
        if isinstance(value, list):
            return [float_bits(val) for val in value]
        return value
    for key in ("raw_data", "entered_effects", "ratio_inputs"):
        actual = saved_row[key]
        if key == "entered_effects":
            # Visiting a previously unused metric creates an empty effect slot.
            actual = [effect for effect in actual if any(value is not None for value in effect["values"].values())]
        assert float_bits(actual) == float_bits(baseline[key]), key
    assert window.open(window.out_path, raise_on_error=True)
    assert raw() == expected
    method = {
        "binary": "binary.random",
        "continuous": "continuous.random",
        "diagnostic": "diagnostic.random",
    }[family]
    record_step(
        "calculator-reopened", window, family=family, raw=raw(), project=window.out_path
    )
    analyze(app, window, method).close()


def grid_entry(app, window):
    from PyQt6.QtTest import QTest
    from rc_metastudio import main_wizard
    from PyQt6 import QtCore, QtWidgets
    # Use the real wizard, clipboard, delegate and workspace history.
    def create():
        dialog = app.activeModalWidget()
        dialog.currentPage().twoarm_proportions_Button.click()
        dialog.button(QtWidgets.QWizard.WizardButton.NextButton).click()
        if isinstance(dialog.currentPage(), main_wizard.ChooseMetricPage):
            dialog.button(QtWidgets.QWizard.WizardButton.NextButton).click()
        dialog.currentPage().outcome_name_LineEdit.setText("Grid entry")
        dialog.button(QtWidgets.QWizard.WizardButton.FinishButton).click()
    QtCore.QTimer.singleShot(100, create)
    window.action_new_dataset.trigger()
    table = window.tableView
    fixture = (ROOT / "tests/e2e/fixtures/hillclimb/binary.tsv").read_text()
    rows = [line.split("\t") for line in fixture.strip().splitlines()]
    rows[0][2] = "bad"
    app.clipboard().setText("\r\n".join("\t".join(row) for row in rows) + "\r\n")
    table.setCurrentIndex(window.model.index(0, window.model.NAME))
    QTest.keyClick(table, QtCore.Qt.Key.Key_V, QtCore.Qt.KeyboardModifier.ControlModifier)
    assert sum(row.include for row in window.model.current_sheet.rows) == 11
    assert "1 cell" in window.statusBar().currentMessage()
    window.undo(); assert not window.model.current_sheet.rows
    window.redo(); assert len(window.model.current_sheet.rows) == 12
    model = window.model
    model.setData(model.index(1, 0), False, QtCore.Qt.ItemDataRole.CheckStateRole)
    model.add_covariate("Dose", "continuous")
    column = table.get_covariate_columns()[-1]
    table.setCurrentIndex(model.index(0, column))
    for value in ("21", "22", "23"):
        QTest.keyClicks(table, value[0])
        editor = next(w for w in table.findChildren(QtWidgets.QLineEdit) if w.isVisible())
        QTest.keyClicks(editor, value[1:])
        QTest.keyClick(editor, QtCore.Qt.Key.Key_Return)
        app.processEvents()
    cov = model.current_sheet.covariates[-1]["id"]
    assert [r.covariate_values.get(cov) for r in model.current_sheet.rows[:3]] == [21, 22, 23]
    window.undo(); window.redo()
    model = window.model
    selection = QtCore.QItemSelection(model.index(0, column), model.index(2, column))
    table.selectionModel().select(selection, QtCore.QItemSelectionModel.SelectionFlag.ClearAndSelect)
    QTest.keyClick(table, QtCore.Qt.Key.Key_D, QtCore.Qt.KeyboardModifier.ControlModifier)
    assert [r.covariate_values.get(cov) for r in model.current_sheet.rows[:3]] == [21]*3
    window.undo(); window.redo()
    model = window.model
    table.selectionModel().select(QtCore.QItemSelection(model.index(0, column), model.index(2, column)), QtCore.QItemSelectionModel.SelectionFlag.ClearAndSelect)
    QTest.keyClick(table, QtCore.Qt.Key.Key_Delete)
    assert all(r.covariate_values.get(cov) is None for r in model.current_sheet.rows[:3])
    window.undo(); window.redo()
    model = window.model
    assert not model.current_sheet.rows[1].include
    assert sum(r.include for r in model.current_sheet.rows) == 10
    record_step("grid-entry-complete", window, included=10, skipped=1)
    import csv
    import io
    block = io.StringIO(newline="")
    writer = csv.writer(block, delimiter="\t", lineterminator="\r\n")
    for row in range(50):
        writer.writerow([f"Extra {row}\tquoted\nname", 2020, 5, 80, 10, 90])
    app.clipboard().setText(block.getvalue())
    table.setCurrentIndex(model.index(12, model.NAME))
    QTest.keyClick(table, QtCore.Qt.Key.Key_V, QtCore.Qt.KeyboardModifier.ControlModifier)
    assert len(model.current_sheet.rows) == 62
    assert model.current_sheet.rows[-1].label == "Extra 49\tquoted\nname"
    window.undo(); assert len(window.model.current_sheet.rows) == 12
    window.redo(); assert len(window.model.current_sheet.rows) == 62
    record_step("grid-entry-large-quoted-paste", window, rows=62)



def hillclimb(app, window):
    """Fixed effort route; setup/cleanup outside task boundaries is not counted."""
    import time
    from rc_metastudio import main_wizard, add_new_dialogs
    from rc_metastudio.diagnostic_metrics_dialog import DiagnosticMetricsDialog
    from rc_metastudio.subgroup_analysis_dialog import SubgroupAnalysisDialog
    from rc_metastudio.publication_bias_dialog import PublicationBiasDialog

    measures = {}
    counts = {}
    task = None

    def event(kind, amount=1, label=""):
        counts[kind] += amount
        record_step("effort", task=task, kind=kind, amount=amount, label=label)

    def click(widget):
        assert widget.isEnabled(), widget.objectName()
        event("clicks", label=widget.objectName() or widget.text())
        if isinstance(widget, QtWidgets.QCheckBox):
            QTest.mouseClick(widget, QtCore.Qt.MouseButton.LeftButton, pos=QtCore.QPoint(8, widget.height() // 2))
        else:
            mouse_click(widget, QtCore.Qt.MouseButton.LeftButton)
        app.processEvents()

    def action(handle):
        assert handle.isEnabled(), handle.text()
        event("clicks", label=handle.text())
        handle.trigger()
        app.processEvents()

    def key(widget, value, modifier=QtCore.Qt.KeyboardModifier.NoModifier):
        event("shortcuts", label=str(value))
        QTest.keyClick(widget, value, modifier)
        app.processEvents()

    def text(widget, value):
        event("typed_characters", len(value), value)
        QTest.keyClicks(widget, value)
        app.processEvents()

    def visible(cls):
        found = [w for w in app.topLevelWidgets() if isinstance(w, cls) and w.isVisible()]
        assert len(found) == 1, (cls, found)
        return found[0]

    def dialog(cls):
        result = visible(cls)
        event("dialogs", label=cls.__name__)
        record_step("hillclimb-dialog", result, task=task)
        return result

    def modal(cls, operate, trigger):
        failures = []
        def callback():
            try:
                operate(dialog(cls))
            except BaseException as error:
                failures.append(error)
                if app.activeModalWidget():
                    app.activeModalWidget().reject()
        QtCore.QTimer.singleShot(0, callback)
        trigger()
        if failures:
            raise failures[0]

    def ok(form):
        click(form.buttonBox.button(QtWidgets.QDialogButtonBox.StandardButton.Ok))

    def begin(name):
        nonlocal counts, task, started
        task = name
        counts = dict(clicks=0, shortcuts=0, typed_characters=0, dialogs=0)
        started = time.perf_counter()

    def end(view, completable=True):
        measures[task] = counts | {"wall_seconds": time.perf_counter() - started,
                                   "completable": completable}
        if view is not None:
            numbers = parsed_numeric_sections(view.results)
            measures[task]["numbers"] = numbers
            measures[task]["texts"] = dict(view.texts)
            record_step("hillclimb-result", view, task=task, measures=measures[task],
                        texts=dict(view.texts), plots=dict(view.results.display_images))
        (EVIDENCE / "measures.json").write_text(json.dumps(measures, indent=2) + "\n")

    def run_analysis(handle, subgroup=False, diagnostic=False, regression=False):
        action(handle)
        if subgroup:
            ok(dialog(SubgroupAnalysisDialog))
        if diagnostic:
            click(dialog(DiagnosticMetricsDialog).btn_ok)
        form = dialog(analysis_setup_dialog.AnalysisSetupDialog)
        if regression:
            for cov, checkbox in form.covs_and_check_boxes:
                if checkbox.isChecked() != (cov["name"] == "Dose"):
                    click(checkbox)
            assert [cov["name"] for cov, checkbox in form.covs_and_check_boxes if checkbox.isChecked()] == ["Dose"]
        ok(form)
        view = visible(results_window.ResultsWindow)
        assert view.results.sections
        return view

    started = 0
    begin("T1")
    def create(wizard):
        click(wizard.currentPage().twoarm_proportions_Button)
        click(wizard.button(QtWidgets.QWizard.WizardButton.NextButton))
        if isinstance(wizard.currentPage(), main_wizard.ChooseMetricPage):
            click(wizard.button(QtWidgets.QWizard.WizardButton.NextButton))
        field = wizard.currentPage().outcome_name_LineEdit
        click(field)
        text(field, "Mortality")
        click(wizard.button(QtWidgets.QWizard.WizardButton.FinishButton))
    modal(main_wizard.MainWizard, create, lambda: action(window.action_new_dataset))
    fixture = ROOT / "tests/e2e/fixtures/hillclimb/binary.tsv"
    app.clipboard().setText(fixture.read_text().rstrip("\n"))
    index = window.model.index(0, window.model.NAME)
    window.tableView.scrollTo(index)
    event("clicks", label="first study cell")
    QTest.mouseClick(window.tableView.viewport(), QtCore.Qt.MouseButton.LeftButton,
                     pos=window.tableView.visualRect(index).center())
    key(window.tableView, QtCore.Qt.Key.Key_V, QtCore.Qt.KeyboardModifier.ControlModifier)
    assert len(window.model.get_studies(False)) == 12
    assert all(s.include for s in window.model.get_studies(False))
    basic = run_analysis(window.action_go)
    end(basic)
    window.out_path = str(WORK / "hillclimb-t1.rcms")
    assert window.save()
    basic.close()

    begin("T2")
    values = json.loads((ROOT / "tests/e2e/fixtures/hillclimb/covariates.json").read_text())
    for name, kind in (("Region", "factor"), ("Dose", "continuous")):
        def add(form):
            click(form.covariate_name_le)
            text(form.covariate_name_le, name)
            if form.datatype_cbo_box.currentData() != kind:
                click(form.datatype_cbo_box)
                key(form.datatype_cbo_box, QtCore.Qt.Key.Key_Down)
                key(form.datatype_cbo_box, QtCore.Qt.Key.Key_Return)
            ok(form)
        modal(add_new_dialogs.AddCovariateDialog, add, lambda: action(window.action_add_covariate))
        column = window.tableView.get_covariate_columns()[-1]
        index = window.model.index(0, column)
        window.tableView.scrollTo(index)
        event("clicks", label=f"{name} first cell")
        QTest.mouseClick(window.tableView.viewport(), QtCore.Qt.MouseButton.LeftButton,
                         pos=window.tableView.visualRect(index).center())
        for value in values[name]:
            value = str(value)
            text(window.tableView, value[0])
            editor = next(w for w in window.tableView.findChildren(QtWidgets.QLineEdit) if w.isVisible())
            text(editor, value[1:])
            key(editor, QtCore.Qt.Key.Key_Return)
            app.processEvents()
        assert [s.covariate_values[next(c["id"] for c in window.model.current_sheet.covariates if c["name"] == name)] for s in window.model.get_studies(False)] == values[name], [s.covariate_values[next(c["id"] for c in window.model.current_sheet.covariates if c["name"] == name)] for s in window.model.get_studies(False)]
    record_step("covariates-entered", window, covariates=[(c["name"], c["data_type"]) for c in window.model.current_sheet.covariates], included=[s.include for s in window.model.get_studies(False)])
    assert all(s.include for s in window.model.get_studies(False))
    subgroup = run_analysis(window.action_subgroup_ma, subgroup=True)
    record_step("hillclimb-subgroup", subgroup, numbers=parsed_numeric_sections(subgroup.results), texts=dict(subgroup.texts))
    subgroup_numbers = parsed_numeric_sections(subgroup.results)
    subgroup.close()
    regression = run_analysis(window.action_meta_regression, regression=True)
    end(regression)
    measures["T2"]["subgroup_numbers"] = subgroup_numbers
    regression.close()

    begin("T3")
    modal(PublicationBiasDialog,
          lambda form: click(form.button_box.button(QtWidgets.QDialogButtonBox.StandardButton.Ok)),
          lambda: action(window.action_publication_bias))
    bias = visible(results_window.ResultsWindow)
    end(bias)
    bias.close()

    # T4 starts from the committed populated diagnostic project, as PLAN specifies.
    window.workspace.mark_saved()
    assert window.open(str(ROOT / "sample_projects/lymph.rcms"), raise_on_error=True)
    begin("T4")
    diagnostic = run_analysis(window.action_go, diagnostic=True)
    end(diagnostic)
    diagnostic.close()

    # Restore T1 input for manuscript export; setup is excluded.
    window.workspace.mark_saved()
    assert window.open(str(WORK / "hillclimb-t1.rcms"), raise_on_error=True)
    basic = analyze(app, window, "binary.random")
    begin("T5")
    def save_picker(*args, **kwargs):
        picker = QtWidgets.QFileDialog(basic)
        picker.setOption(QtWidgets.QFileDialog.Option.DontUseNativeDialog)
        picker.setAcceptMode(QtWidgets.QFileDialog.AcceptMode.AcceptSave)
        picker.setDirectory(str(WORK))
        suffix = next(formats)
        picker.setNameFilter("Report (*." + suffix + ")")
        def save(form):
            field = form.findChild(QtWidgets.QLineEdit, "fileNameEdit")
            click(field)
            text(field, "hillclimb-export." + suffix)
            click(form.findChild(QtWidgets.QDialogButtonBox).button(QtWidgets.QDialogButtonBox.StandardButton.Save))
        modal(QtWidgets.QFileDialog, save, picker.exec)
        return picker.selectedFiles()[0], picker.selectedNameFilter()
    formats = iter(("pdf", "docx"))
    with patch.object(QtWidgets.QFileDialog, "getSaveFileName", side_effect=save_picker):
        action(basic.export_report_action)
        action(basic.export_report_action)
    assert (WORK / "hillclimb-export.docx").stat().st_size
    assert (WORK / "hillclimb-export.pdf").read_bytes().startswith(b"%PDF")
    end(basic, True)
    basic.close()


base_path_patch = patch(
    "rc_metastudio.settings.get_base_path", return_value=WORK.as_posix()
)
base_path_patch.start()
try:
    _startup_phases["composition:start"] = time.time()
    app, window = automation.start_automation(
        phase_callback=lambda phase: _startup_phases.update({phase: time.time()})
    )
    assert "rpy2.robjects" not in sys.modules, "R imported before the window appeared"
    assert window.isVisible()
    assert not window.action_go.isEnabled()
    from rc_metastudio import r_bridge, r_backend
    try:
        r_bridge.execute_r_string("1 + 1")
    except r_backend.AnalysisBackendUnavailableError as error:
        assert str(error) == "R is still starting", str(error)
    else:
        raise AssertionError("pre-ready R request succeeded")
    app.processEvents()
    assert r_backend.is_ready()
    if EVIDENCE is not None and os.environ.get("RCMS_DRIVE_STARTED"):
        import time
        app.processEvents()
        assert window.isVisible()
        (EVIDENCE / "startup.json").write_text(json.dumps({"seconds": _startup_phases["window:visible"] - float(os.environ["RCMS_DRIVE_STARTED"]), "visible": True, "platform": app.platformName()}))
    if _startup_profile is not None and EVIDENCE is not None:
        _startup_phases["r:ready"] = time.time()
        _startup_profile.disable()
        _startup_profile.dump_stats(str(EVIDENCE / "startup.pstats"))
        (EVIDENCE / "startup-phases.json").write_text(json.dumps(_startup_phases, indent=2))
except BaseException:
    base_path_patch.stop()
    sys.excepthook = sys.__excepthook__
    raise
sys.excepthook = sys.__excepthook__
errors = []


def dismiss_unexpected_messages():
    for widget in app.topLevelWidgets():
        if isinstance(widget, QtWidgets.QMessageBox) and widget.isVisible() and not widget.property("verificationExpectedMessage"):
            errors.append(widget.text())
            widget.setProperty("verificationExpectedMessage", True)
            record_step("unexpected-dialog", widget, message=widget.text())
            print("UNEXPECTED DIALOG:", widget.text(), file=sys.stderr, flush=True)
            widget.reject()


watchdog = QtCore.QTimer()
watchdog.timeout.connect(dismiss_unexpected_messages)
watchdog.start(100)
try:
    scenario = sys.argv[1]
    if EVIDENCE is not None or scenario == "doctor":
        doctor(app, window)
    if scenario == "doctor":
        pass
    elif scenario == "sheets-baseline":
        from tests.e2e.sheets_baseline import run

        run(app, window, record_step, analyze)
    elif scenario == "grid-entry":
        grid_entry(app, window)
    elif scenario == "hillclimb":
        hillclimb(app, window)
    elif scenario.startswith("audit-round3-"):
        sys.path.insert(0, str(ROOT / ".agents/skills/verify-rcms/scripts"))
        from tests.e2e.audit_round3 import run

        run(app, window, WORK, record_step, scenario)
    elif scenario == "analysis-dialog":
        analysis_dialog(app, window)
    elif scenario == "review1-design":
        sys.path.insert(0, str(ROOT / ".agents/skills/verify-rcms/scripts"))
        from tests.e2e.review1_design import run

        run(app, window, WORK, ROOT, record_step, analyze)
    elif scenario == "review-workflows":
        from tests.e2e.review_workflows import run

        run(app, window, ROOT, record_step, analyze)
    elif scenario in ("ux-improvements", "results-window"):
        from tests.e2e.ux_workflows import run

        run(app, window, WORK, ROOT, record_step, analyze)
    elif scenario == "manuscript-export":
        from tests.e2e.manuscript_export import run
        run(app, window, WORK, ROOT, record_step, analyze)
    elif scenario in ("amino", "continuous", "lymph"):
        sample_analysis(app, window, scenario)
    elif scenario == "rm5-import":
        from tests.e2e.rm5_import import run
        run(app, window, ROOT, WORK, record_step)
    elif scenario == "sheets-persist":
        sheets_persist(app, window)
    elif scenario == "sheets-edges":
        sheets_edges(app, window)
    elif scenario == "sheets-swap":
        sheets_swap(app, window)
    elif scenario == "sheets-legacy":
        sheets_legacy(app, window)
    elif scenario in ("sheets", "project-structure"):
        sheets(app, window)
    elif scenario == "editing":
        editing(app, window)
    elif scenario == "plot-edit":
        plot_edit(app, window)
    elif scenario == "small-study-effects":
        small_study_effects(app, window)
    elif scenario.startswith(
        (
            "options-",
            "plot-workflow-",
            "calculator-mode-",
            "calculator-grid-precision",
            "calculator-ratio-precision",
            "project-",
            "analysis-sequential-",
        )
    ):
        sys.path.insert(0, str(ROOT / ".agents/skills/verify-rcms/scripts"))
        from advanced_workflows import run

        run(app, window, scenario, record_step)
    elif scenario.startswith("calculator-"):
        calculator(app, window, scenario.removeprefix("calculator-"))
    elif scenario == "xlsx-import":
        xlsx_import(app, window)
    elif scenario == "csv":
        csv_import(app, window)
    elif scenario.startswith("ratios-"):
        sys.path.insert(0, str(ROOT / ".agents/skills/verify-rcms/scripts"))
        from ratio_workflows import run

        run(app, window, scenario, record_step)
    elif scenario.startswith(("analysis-", "small-study-")):
        sys.path.insert(0, str(ROOT / ".agents/skills/verify-rcms/scripts"))
        from analysis_matrix import run

        run(app, window, scenario, record_step)
    else:
        raise ValueError(f"Unknown desktop scenario: {scenario}")
    assert not errors, errors
    record_step("workflow-passed", window, scenario=scenario)
except BaseException:
    traceback.print_exc()
    raise
finally:
    window.workspace.mark_saved()
    watchdog.stop()
    for widget in app.topLevelWidgets():
        widget.close()
    automation.dispose_qobjects(app, app.topLevelWidgets())
    app.quit()
    base_path_patch.stop()
print("WORKFLOW PASSED")
