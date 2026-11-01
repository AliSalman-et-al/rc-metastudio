"""Drive the all-outcome summary through the Analysis menu."""
from PyQt6 import QtWidgets
from docx import Document


def run(app, window, work, root, record, analyze):
    assert window.open(str(root / 'sample_projects/amino.rcms'), raise_on_error=True)
    import copy
    from tests.analysis_regression.golden.support.headless_analysis import load_dataset_model
    diagnostic = copy.deepcopy(load_dataset_model(root / 'sample_projects/lymph.rcms').dataset.sheets[0])
    empty = copy.deepcopy(window.model.dataset.sheets[0])
    from rc_metastudio.project_domain import SheetId
    import uuid
    empty.id = SheetId(uuid.uuid4().hex)
    diagnostic.id = SheetId(uuid.uuid4().hex)
    empty.name = 'Empty outcome'
    empty.rows = []
    window.model.dataset.sheets.extend((empty, diagnostic))
    window.data_dirtied()
    window.out_path = str(work / 'summary-fixture.rcms')
    assert window.save()
    action = next((a for a in window.menuAnalysis.actions()
                   if a.text() == 'Summary of findings'), None)
    assert action is not None, 'Analysis menu lacks Summary of findings'
    action.trigger()
    app.processEvents()
    dialog = window.summary_of_findings_dialog
    assert dialog.table.rowCount() == 4
    for row in range(2):
        assert dialog.table.item(row, 1).text() != '0'
        assert 'CI' in dialog.table.item(row, 3).text()
        assert dialog.table.cellWidget(row, 4).currentText() == 'Not rated'
    assert 'Too little data' in dialog.table.item(2, 3).text()
    record('summary-diagnostic', dialog, effect=dialog.table.item(3, 3).text())
    screens = root / 'docs/ux-hill-climb/screens'
    dialog.grab().save(str(screens / 'slice9-sof-light.png'))
    from PyQt6 import QtGui
    original = app.palette()
    palette = QtGui.QPalette(original)
    for role in (QtGui.QPalette.ColorRole.Text, QtGui.QPalette.ColorRole.ButtonText, QtGui.QPalette.ColorRole.WindowText):
        palette.setColor(role, QtGui.QColor('#eeeeee'))
    for role in (QtGui.QPalette.ColorRole.Window, QtGui.QPalette.ColorRole.Button, QtGui.QPalette.ColorRole.Base):
        palette.setColor(role, QtGui.QColor('#202020'))
    app.setPalette(palette)
    app.processEvents()
    dialog.grab().save(str(screens / 'slice9-sof-dark.png'))
    app.setPalette(original)
    report = dialog.report()
    from rc_metastudio.manuscript_report import write_docx, write_pdf
    write_docx(report, work / 'summary.docx')
    write_pdf(report, work / 'summary.pdf')
    document = Document(work / 'summary.docx')
    assert len(document.tables[0].rows) == 5
    assert document.tables[0].cell(1, 4).text == 'Not rated'
    record('summary-of-findings', dialog, rows=report.sections[0].tables[0].rows)
    import subprocess
    pdf_text = subprocess.check_output(['pdftotext', str(work / 'summary.pdf'), '-'], text=True)
    import unicodedata
    pdf_text = unicodedata.normalize('NFKC', pdf_text)
    assert 'Summary of findings' in pdf_text and 'Not rated' in pdf_text
    dialog.table.cellWidget(0, 4).setCurrentText('Moderate')
    dialog.table.cellWidget(0, 5).setText('Reviewer judgement')
    dialog.table.cellWidget(0, 5).editingFinished.emit()
    dialog.close()
    assert window.save()
    assert window.open(str(work / 'summary-fixture.rcms'), raise_on_error=True)
    saved = window.model.dataset.sheets[0].settings['summary_of_findings']
    assert saved == {'certainty': 'Moderate', 'comments': 'Reviewer judgement'}
    import zipfile, json
    with zipfile.ZipFile(work / 'summary-fixture.rcms') as archive:
        project = json.loads(archive.read('project.json'))
    assert all('results' not in sheet for sheet in project['sheets'])
