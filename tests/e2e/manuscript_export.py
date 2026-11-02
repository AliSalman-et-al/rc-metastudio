from typing import Any, Callable, cast
import sys
import gc
"""Reopen real manuscript exports and compare with the executed result."""
from unittest.mock import patch
import subprocess
import unicodedata
from rc_metastudio import __version__
from docx import Document
from PyQt6 import QtWidgets, QtCore, QtTest, sip
from rc_metastudio import automation
from tests.python.gui.test_types import mouse_click, required


def run(app, window, work, root, record, analyze, *, direct_write=False):
    """Compare the same writers/assertions with and without a real save dialog."""
    from collections import Counter
    from contextlib import ExitStack
    import weakref
    from rc_metastudio import manuscript_report

    mode = "direct" if direct_write else "dialog"
    writes = Counter()

    def checked_writer(extension, writer):
        def write(report, path):
            resources = []

            def observed(factory):
                def create(*args, **kwargs):
                    value = factory(*args, **kwargs)
                    resources.append((factory.__name__, weakref.ref(value)))
                    return value
                return create

            state = dict(mode=mode, format=extension, path=str(path), number=sum(writes.values()) + 1)
            print("Manuscript writer: " + repr(dict(state, phase="enter")), file=sys.stderr, flush=True)
            gc_enabled = gc.isenabled()
            gc.disable()
            try:
                with ExitStack() as stack:
                    for name in ("QPdfWriter", "QPainter", "QTextDocument"):
                        factory = getattr(manuscript_report, name)
                        stack.enter_context(patch.object(manuscript_report, name, new=observed(factory)))
                    writer(report, path)
                # Reference-counted Qt resources must go away on return, without
                # cyclic collection concealing a retained writer or paint device.
                alive = [name for name, reference in resources if reference() is not None]
            finally:
                if gc_enabled:
                    gc.enable()
            created = dict(Counter(name for name, _reference in resources))
            print("Manuscript writer: " + repr(dict(state, phase="returned", created=created, alive=alive)), file=sys.stderr, flush=True)
            assert not alive, (state, "writer resources retained after return", alive)
            assert pathlib.Path(path).is_file(), path
            writes[extension] += 1
            record("manuscript-writer-released", **state, created=created, alive=alive)
        return write

    import pathlib
    with patch.object(manuscript_report, "write_pdf", new=checked_writer("pdf", manuscript_report.write_pdf)), \
         patch.object(manuscript_report, "write_docx", new=checked_writer("docx", manuscript_report.write_docx)):
        _run(app, window, work, root, record, analyze, direct_write=direct_write)
    assert writes == {"pdf": 20, "docx": 8}, (mode, writes)
    record("manuscript-export-mode-complete", mode=mode, writes=dict(writes))


def _run(app, window, work, root, record, analyze, *, direct_write):
    pending_close = None
    previous_dialog: QtWidgets.QFileDialog | None = None
    picker_number = 0
    operation = "setup"
    mode = "direct" if direct_write else "dialog"

    def export_report(view, path, suffix, picker=None):
        if direct_write:
            from rc_metastudio.manuscript_report import ManuscriptReport, write_docx, write_pdf
            (write_pdf if suffix == "pdf" else write_docx)(ManuscriptReport.from_result(view.results), path)
        elif picker is None:
            with patch.object(QtWidgets.QFileDialog, 'getSaveFileName', return_value=(str(path), suffix)):
                view.export_report_action.trigger()
        else:
            with patch.object(QtWidgets.QFileDialog, 'getSaveFileName', side_effect=picker):
                view.export_report_action.trigger()

    def lifetime_marker(phase, view):
        deleted = sip.isdeleted(view)
        state = {"phase": phase, "view": hex(id(view)), "deleted": deleted,
                 "mode": mode, "operation": operation, "picker_number": picker_number,
                 "previous_dialog_deleted": None if previous_dialog is None else sip.isdeleted(previous_dialog)}
        if pending_close is not None:
            state["previous_result_deleted"] = sip.isdeleted(pending_close)
        if not deleted:
            state.update(visible=view.isVisible(), refit_pending=view._viewport_refit_pending,
                         scene_deleted=sip.isdeleted(view.scene),
                         viewport_deleted=sip.isdeleted(view.graphics_view.viewport()))
        print("Manuscript lifetime: " + repr(state), file=sys.stderr, flush=True)

    def close_result(view):
        lifetime_marker("close-before", view)
        view.close()
        automation.dispose_qobjects(app, (view,))
        assert sip.isdeleted(view)
        gc.collect()
        app.processEvents()
        lifetime_marker("close-after", view)

    # Derive boundary cases from real legacy bytes, then open through production IO.
    import copy, io, json, zipfile
    legacy = subprocess.check_output(['git', 'show', '322ede26:sample_projects/amino.rcms'], cwd=root)
    with zipfile.ZipFile(io.BytesIO(legacy)) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    project = json.loads(members['project.json'])
    dataset = project['dataset']
    blank = copy.deepcopy(dataset['studies'][-1])
    blank.update(name='', year=None, sample_size=None, include=False, manually_excluded=False, notes='', covariates={})
    for unit in blank['analysis_units']:
        unit['entered_effects'] = {}
        for arm in unit['groups']:
            arm['raw_data'] = [''] * len(arm['raw_data'])
    dataset['covariates'].append({'name': 'Only covariate', 'data_type': 1, 'stable_id': 'only-covariate'})
    variants: list[dict[str, Any]] = [blank, dict(blank, name='Named pending'), dict(blank, covariates={'Only covariate': 7}), dict(blank, year=2001), dict(blank, manually_excluded=True)]
    for i, row in enumerate(variants):
        row = copy.deepcopy(row)
        row['id'] = 10000 + i
        dataset['studies'].append(row)
    for row in dataset['studies']:
        row['covariates'].setdefault('Only covariate', None)
    members['project.json'] = json.dumps(project).encode()
    import hashlib
    manifest = json.loads(members['manifest.json'])
    manifest['members']['project.json'] = {'size': len(members['project.json']), 'sha256': hashlib.sha256(members['project.json']).hexdigest()}
    members['manifest.json'] = json.dumps(manifest).encode()
    fixture = work / 'legacy-row-boundaries.rcms'
    with zipfile.ZipFile(fixture, 'w') as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    assert window.open(str(fixture), raise_on_error=True)
    retained = window.model.get_studies(only_if_included=False)
    assert any(row.label == 'Named pending' for row in retained)
    assert any(row.year == 2001 for row in retained)
    assert any(row.manually_excluded for row in retained)
    assert any(7 in row.covariate_values.values() for row in retained)
    assert len(retained) == 23, len(retained)
    record('legacy-row-boundaries', window, retained=len(retained), input=str(fixture))
    with zipfile.ZipFile(root / 'sample_projects/amino.rcms') as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    project = json.loads(members['project.json'])
    row = project['sheets'][0]['rows'][0]
    for values in row['raw_data'].values():
        values[0] = 0
    row['entered_effects'] = []
    members['project.json'] = json.dumps(project).encode()
    manifest = json.loads(members['manifest.json'])
    manifest['members']['project.json'] = {'size': len(members['project.json']), 'sha256': hashlib.sha256(members['project.json']).hexdigest()}
    members['manifest.json'] = json.dumps(manifest).encode()
    zero_fixture = work / 'double-zero.rcms'
    with zipfile.ZipFile(zero_fixture, 'w') as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    window.workspace.mark_saved()
    assert window.open(str(zero_fixture), raise_on_error=True)
    print("Manuscript analysis: double-zero binary.random", file=sys.stderr, flush=True)
    view = analyze(app, window, "binary.random")
    assert "3 studies with no events in either arm were retained with this correction." in view.manuscript.methods, view.manuscript.methods
    record('manuscript-double-zero', view, methods=view.manuscript.methods)
    close_result(view)
    # Real result Methods must describe where correction enters each estimator.
    for method in ("binary.fixed.mh", "binary.fixed.peto"):
        window.workspace.mark_saved()
        assert window.open(str(root / 'sample_projects/amino.rcms'), raise_on_error=True)
        print(f"Manuscript analysis: {method}", file=sys.stderr, flush=True)
        view = analyze(app, window, method)
        expected = "A continuity correction of 0.5 was applied to individual study estimates only; the pooled " + ("Mantel-Haenszel" if method.endswith("mh") else "Peto") + " estimate uses the original counts."
        assert expected in view.manuscript.methods, view.manuscript.methods
        (work / (method + '-prose.txt')).write_text(view.manuscript.methods, encoding='utf-8')
        for suffix in ('pdf', 'docx'):
            path = work / (method + '.' + suffix)
            export_report(view, path, suffix)
            if suffix == 'pdf':
                text = subprocess.check_output(['pdftotext', '-enc', 'UTF-8', str(path), '-'], text=True, encoding='utf-8')
            else:
                doc = Document(path)
                text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            assert ' '.join(expected.split()) in ' '.join(text.split())
            import re
            assert not re.search(r"\b(?:NaN|Inf)\b", text), text
        record('manuscript-correction-' + method, view, methods=view.manuscript.methods)
        close_result(view)
    for sample, method in (("amino", "binary.random"), ("continuous", "continuous.random"), ("lymph", "diagnostic.reitsma"), ("Pignon2000-HR", "time_to_event.random")):
        window.workspace.mark_saved()
        assert window.open(str(root / 'sample_projects' / (sample + '.rcms')), raise_on_error=True)
        print(f"Manuscript analysis: {method}", file=sys.stderr, flush=True)
        view = analyze(app, window, method)
        def picker(*args, **kwargs):
            nonlocal previous_dialog, picker_number
            assert previous_dialog is None or sip.isdeleted(previous_dialog)
            picker_number += 1
            assert args[0] is view and not sip.isdeleted(view) and view.isVisible()
            lifetime_marker(f"{sample}-{suffix}-picker-create", view)
            record('manuscript-picker-before', view, sample=sample, suffix=suffix)
            dialog = QtWidgets.QFileDialog(view)
            dialog.setOption(QtWidgets.QFileDialog.Option.DontUseNativeDialog)
            dialog.setAcceptMode(QtWidgets.QFileDialog.AcceptMode.AcceptSave)
            dialog.setNameFilter('Report (*.' + suffix + ')')
            dialog.setDirectory(str(work))
            dialog.selectFile(str(path))
            models = dialog.findChildren(QtCore.QAbstractItemModel)
            model_names = [required(model.metaObject(), "file-dialog model type").className() for model in models]
            print("Manuscript dialog: " + repr(dict(operation=operation, picker_number=picker_number,
                  phase="created", directory=str(work), models=model_names)), file=sys.stderr, flush=True)
            picker_errors = []
            def save():
                try:
                    assert cast(Callable[..., bool], QtTest.QTest.qWaitForWindowExposed)(dialog, 5000)
                    gc.collect()
                    app.processEvents()
                    lifetime_marker(f"{sample}-{suffix}-picker-visible", view)
                    assert not sip.isdeleted(view) and view.isVisible()
                    from pathlib import Path
                    assert Path(dialog.selectedFiles()[0]).resolve() == path.resolve()
                    record('manuscript-picker-visible', dialog, sample=sample, suffix=suffix)
                    box = required(dialog.findChild(QtWidgets.QDialogButtonBox))
                    button = required(box.button(QtWidgets.QDialogButtonBox.StandardButton.Save))
                    assert button.isEnabled()
                    mouse_click(button, QtCore.Qt.MouseButton.LeftButton)
                except BaseException as error:
                    picker_errors.append(error)
                    dialog.reject()
            QtCore.QTimer.singleShot(0, save)
            try:
                lifetime_marker(f"{sample}-{suffix}-picker-exec-enter", view)
                accepted = dialog.exec()
                lifetime_marker(f"{sample}-{suffix}-picker-exec-return", view)
                assert not picker_errors, picker_errors
                assert accepted
                selected = dialog.selectedFiles()[0], dialog.selectedNameFilter()
            finally:
                # Explicitly finish this dialog's deferred destruction before
                # the next writer or nested event loop. This also checks models.
                dialog.hide()
                dialog.deleteLater()
                for _ in range(2):
                    app.sendPostedEvents(None, QtCore.QEvent.Type.DeferredDelete)
                    app.processEvents()
                destroyed = sip.isdeleted(dialog)
                models_destroyed = [sip.isdeleted(model) for model in models]
                print("Manuscript dialog: " + repr(dict(operation=operation, picker_number=picker_number,
                      phase="deleted", dialog_deleted=destroyed, models=model_names,
                      models_deleted=models_destroyed)), file=sys.stderr, flush=True)
                assert destroyed and all(models_destroyed), (operation, destroyed, models_destroyed)
                previous_dialog = dialog
            assert sip.isdeleted(dialog)
            assert not sip.isdeleted(view) and view.isVisible()
            record('manuscript-picker-after', view, sample=sample, suffix=suffix)
            return selected
        for suffix in ('pdf', 'docx'):
            operation = f"{sample}-{suffix}"
            path = work / (sample + '.' + suffix)
            export_report(view, path, suffix, picker)
            assert path.exists()
            if suffix == 'docx':
                doc = Document(path)
                assert doc.paragraphs[0].text == view.manuscript.title
                assert 'Table 1.' in '\n'.join(p.text for p in doc.paragraphs)
                if sample == 'amino':
                    assert doc.tables[0].rows[0].cells[0].text == 'Statistic'
                assert len(doc.inline_shapes) >= len(view.results.images)
                assert all(round(s.width / 36000) in (85, 180) for s in doc.inline_shapes)
                text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            else:
                text = subprocess.check_output(['pdftotext', '-enc', 'UTF-8', '-layout', str(path), '-'], text=True, encoding='utf-8')
                from xml.etree import ElementTree as ET
                bbox = subprocess.check_output(['pdftotext', '-enc', 'UTF-8', '-bbox', str(path), '-'])
                (work / (sample + '-bbox.html')).write_bytes(bbox)
                pages = ET.fromstring(bbox).findall('.//{*}page')
                tables = [word for page in pages for word in page.findall('.//{*}word') if word.text == 'Table']
                figures = [(page, word) for page in pages for word in page.findall('.//{*}word') if word.text == 'Figure']
                assert tables and figures
                table_height = float(tables[0].attrib['yMax']) - float(tables[0].attrib['yMin'])
                for page, word in figures:
                    figure_height = float(word.attrib['yMax']) - float(word.attrib['yMin'])
                    assert abs(figure_height - table_height) < 1, (sample, figure_height, table_height)
                    assert 0 <= float(word.attrib['yMin']) < float(word.attrib['yMax']) <= float(page.attrib['height'])

            assert 'Methods' in text and 'RC MetaStudio' in ' '.join(text.split())
            normalized = ' '.join(unicodedata.normalize('NFKC', text).split())
            for prose in (view.manuscript.methods, view.manuscript.results):
                assert ' '.join(unicodedata.normalize('NFKC', prose).split()) in normalized, (sample, suffix, prose, normalized)
            assert view.manuscript_methods.toPlainText() == view.manuscript.methods
            assert view.manuscript_results.toPlainText() == view.manuscript.results
            assert all(value not in text for value in ('executed settings', 'Project:', '.rcms'))
            assert 'Supplementary details' in text
            assert text.index('Supplementary details') > text.rindex('Figure ')
            if sample == 'amino':
                assert view.manuscript.title == 'clinical failure: random-effects meta-analysis (odds ratio)'
                assert 'The pooled odds ratio was 0.770 (95% CI 0.485 to 1.222; p = 0.267).' in view.manuscript.results
                assert view.manuscript.results == '19 studies were included (1097 participants in Intervention, 1080 in Control). The pooled odds ratio was 0.770 (95% CI 0.485 to 1.222; p = 0.267). Heterogeneity: Q = 33.360 (df = 18; p = 0.015), I² = 46.0%, τ² = 0.378. The 95% prediction interval was 0.212 to 2.799.', view.manuscript.results
                assert view.manuscript.methods == f"Odds ratios were pooled using a random-effects model with the DerSimonian–Laird estimator of between-study variance. Confidence intervals (95%) used the normal approximation. A continuity correction of 0.5 was applied to studies with any zero cell. 2 studies with no events in either arm were retained with this correction. Heterogeneity was assessed with Cochran's Q and I². Analyses were performed in RC MetaStudio {__version__}. R packages loaded for this analysis: metafor 5.0.1, meta 8.5-0."
                assert '-0.262' in normalized
            record('manuscript-export-' + sample + '-' + suffix, view, exported=str(path))
        (work / (sample + '-prose.txt')).write_text('Methods\n' + view.manuscript.methods + '\n\nResults\n' + view.manuscript.results, encoding='utf-8')
        item = view.nav_tree.findItems('Manuscript', __import__('PyQt6.QtCore', fromlist=['Qt']).Qt.MatchFlag.MatchExactly)[0]
        view.nav_tree.setCurrentItem(item)
        record('manuscript-panel-' + sample, view)
        result = view.results
        close_result(view)
        from rc_metastudio.results_window import ResultsWindow
        for cycle in range(3):
            # Leave a real close/deferred-delete pending across the next save
            # dialog, as opposed to only testing explicit harness disposal.
            previous = ResultsWindow(result, parent=window)
            previous.show()
            previous.resize(previous.width() + 1, previous.height())
            lifetime_marker(f"{sample}-stress-{cycle}-previous-close", previous)
            previous.close()
            pending_close = previous
            lifetime_marker(f"{sample}-stress-{cycle}-previous-closed", previous)
            assert previous.graphics_view.scene() is None, "Closing Results must detach the viewport before deferred deletion"
            view = ResultsWindow(result, parent=window)
            view.show()
            view.resize(view.width() + 1, view.height())
            gc.collect()
            path = work / f"stress-{sample}-{cycle}.pdf"
            suffix = "pdf"
            operation = f"{sample}-stress-{cycle}-pdf"
            lifetime_marker(f"{sample}-stress-{cycle}-export", view)
            export_report(view, path, suffix, picker)
            assert path.read_bytes().startswith(b"%PDF")
            app.processEvents()
            assert sip.isdeleted(previous)
            assert sip.isdeleted(previous.scene)
            assert all(sip.isdeleted(item) for item in previous._layout_items)
            record("manuscript-lifetime-stress", view, sample=sample, cycle=cycle,
                   previous_deleted=True, exported=str(path))
            close_result(view)
            pending_close = None
            del previous
            gc.collect()
            app.processEvents()

    sys.path.insert(0, str(root / ".agents/skills/verify-rcms/scripts"))
    from tests.e2e.audit_round3 import analyze_existing
    for workflow in ('subgroup', 'meta-regression'):
        window.workspace.mark_saved()
        assert window.open(str(root / 'sample_projects/amino.rcms'), raise_on_error=True)
        studies = window.model.get_studies()
        window.model.add_covariate('Group', 'factor', {s.label: 'A' if i % 2 else 'B' for i, s in enumerate(studies)})
        window.model.add_covariate('Dose', 'continuous', {s.label: float(i + 1) for i, s in enumerate(studies)})
        window.model.reset_model()
        window.enable_menu_options_that_require_dataset()
        view = analyze_existing(app, window, workflow, 'binary.random', record)
        from rc_metastudio.manuscript_report import write_docx, write_pdf
        write_docx(view.manuscript, work / (workflow + '.docx'))
        write_pdf(view.manuscript, work / (workflow + '.pdf'))
        assert workflow.split('-')[0].lower() in view.manuscript.methods.lower()
        prose = view.manuscript.results
        if workflow == 'subgroup':
            assert 'Subgroup A: estimate' in prose and 'Subgroup B: estimate' in prose
            assert 'Between-group test:' in prose
        else:
            assert 'coefficient' in prose and '95% CI' in prose and '; p = ' in prose, (prose, [(s.title, s.tables) for s in view.results.sections if s.kind == 'text'])
            assert 'Residual heterogeneity:' in prose
            assert any(t.title == 'Regression coefficients' for section in view.manuscript.sections for t in section.tables)
        exported = Document(work / (workflow + '.docx'))
        captions = [p.text for p in exported.paragraphs]
        assert any(p.startswith('Table ') for p in captions)
        if 'Supplementary details' in captions:
            assert captions.index('Supplementary details') > max(i for i, p in enumerate(captions) if p.startswith('Figure '))
        (work / (workflow + '-prose.txt')).write_text('Methods\n' + view.manuscript.methods + '\n\nResults\n' + view.manuscript.results, encoding='utf-8')
        record('manuscript-' + workflow, view)
        close_result(view)

    from PyQt6.QtCore import QTimer
    from rc_metastudio.publication_bias_dialog import PublicationBiasDialog
    from rc_metastudio.results_window import ResultsWindow
    window.workspace.mark_saved()
    assert window.open(str(root / 'sample_projects/amino.rcms'), raise_on_error=True)
    errors = []
    def run_small_study():
        dialog = app.activeModalWidget()
        try:
            assert isinstance(dialog, PublicationBiasDialog)
            dialog.run()
            assert not dialog.failure_label.isVisible(), dialog.failure_label.text()
        except Exception as error:
            errors.append(error)
        finally:
            dialog.reject()
    QTimer.singleShot(0, run_small_study)
    window.action_publication_bias.trigger()
    assert not errors, errors
    view = next(w for w in app.topLevelWidgets() if isinstance(w, ResultsWindow) and w.isVisible())
    assert 'Harbord' in view.manuscript.methods or 'Rücker' in view.manuscript.methods
    assert not any(word in view.manuscript.methods.lower() for word in ('evidence', 'bias present', 'bias absent'))
    app.processEvents()
    app.processEvents()
    assert view._figure_zoom_percent >= 50, view._figure_zoom_percent
    assert all(item.textWidth() >= view.graphics_view.viewport().width() * .9 for item in view._wrapped_text_items)
    record('manuscript-small-study-methods', view, methods=view.manuscript.methods,
           settled_zoom_percent=view._figure_zoom_percent)
    close_result(view)
