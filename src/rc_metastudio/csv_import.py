# SPDX-FileCopyrightText: 2026 Ali Salman and RC MetaStudio contributors
# SPDX-License-Identifier: GPL-3.0-or-later
"""Parse CSV files into the workspace import contract."""

from __future__ import annotations

import csv
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import TypedDict

from rc_metastudio import tabular_data


class CsvImportError(ValueError):
    """A CSV file cannot satisfy the workspace import contract."""


class CsvImportPayload(TypedDict):
    """Mutable representation retained by the wizard and undo command."""

    headers: list[str]
    data: list[list[str]]
    expected_headers: list[str]
    covariate_names: list[str]
    covariate_types: list[str]


@dataclass(frozen=True, slots=True)
class CsvImportResult:
    """Normalized, validated rows ready for preview and workspace import."""

    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    expected_headers: tuple[str, ...]
    covariate_names: tuple[str, ...]
    covariate_types: tuple[str, ...]

    def to_payload(self) -> CsvImportPayload:
        return {
            "headers": list(self.headers),
            "data": [list(row) for row in self.rows],
            "expected_headers": list(self.expected_headers),
            "covariate_names": list(self.covariate_names),
            "covariate_types": list(self.covariate_types),
        }


def parse_csv(
    path: str | Path,
    *,
    expected_headers: list[str] | tuple[str, ...],
    has_headers: bool,
    from_excel: bool,
    delimiter: str = ",",
    quotechar: str = '"',
    year_column: int = 1,
    optional_headers: Mapping[str, str] | None = None,
    column_mapping: tuple[int, ...] | None = None,
    header_aliases: Mapping[str, str] | None = None,
    worksheet: str | None = None,
    notices: list[str] | None = None,
) -> CsvImportResult:
    """Read, normalize, and validate one CSV file."""
    suffix = Path(path).suffix.lower()
    if suffix == ".xls":
        raise CsvImportError("Save this workbook as .xlsx before importing it.")
    if suffix == ".xlsx":
        headers, rows, _names, _selected = read_xlsx(path, worksheet, year_column=(column_mapping[year_column] if column_mapping else year_column), numeric_columns=set(column_mapping[1:] if column_mapping else range(1, len(expected_headers))), notices=notices)
        has_headers = True
    else:
        with Path(path).open(encoding="utf-8-sig", newline="") as stream:
            reader = (
                csv.reader(stream, dialect="excel")
                if from_excel
                else csv.reader(stream, delimiter=delimiter, quotechar=quotechar)
            )
            headers = next(reader, []) if has_headers else []
            rows = list(reader)
    if header_aliases:
        headers = [header_aliases.get(header.casefold(), header) for header in headers]

    if column_mapping is not None:
        if not has_headers:
            raise CsvImportError("Column mapping requires a header row.")
        if len(column_mapping) != len(expected_headers) or len(set(column_mapping)) != len(column_mapping):
            raise CsvImportError("Map each required field to a different CSV column.")
        if any(index < 0 or index >= len(headers) for index in column_mapping):
            raise CsvImportError("Select a CSV column for every required field.")
        indices = list(column_mapping) + [index for index in range(len(headers)) if index not in column_mapping]
        rows = [[row[index] if index < len(row) else "" for index in indices] for row in rows]
        headers = list(expected_headers) + [headers[index] for index in indices[len(expected_headers):]]

    if has_headers:
        if headers:
            _validate_headers(headers, expected_headers)
        elif rows:
            raise CsvImportError("CSV file is missing the required header row.")
    else:
        _validate_minimum_width(rows, len(expected_headers))
    normalized_rows = normalize_import_rows(rows, minimum_width=len(headers))
    if headers:
        width = len(normalized_rows[0]) if normalized_rows else len(headers)
        headers = headers + [""] * (width - len(headers))

    if optional_headers:
        base_width = len(expected_headers)
        width = len(normalized_rows[0]) if normalized_rows else len(headers)
        optional_indices = {}
        for name in optional_headers:
            matches = [i for i, header in enumerate(headers) if header == name]
            if len(matches) > 1:
                raise CsvImportError(f"CSV contains duplicate input column: {name}.")
            optional_indices[name] = matches[0] if matches else None
        covariate_indices = [i for i in range(base_width, width) if i not in optional_indices.values()]
        expanded_rows = []
        for row in normalized_rows:
            values = row[:base_width]
            for name, default in optional_headers.items():
                index = optional_indices[name]
                values.append(default if index is None else row[index])
            expanded_rows.append(values + [row[i] for i in covariate_indices])
        normalized_rows = expanded_rows
        if has_headers:
            headers = list(expected_headers) + list(optional_headers) + [headers[i] for i in covariate_indices]
        expected_headers = list(expected_headers) + list(optional_headers)

    _validate_years(normalized_rows, year_column)
    covariate_names, covariate_types = _infer_covariates(
        normalized_rows,
        headers=headers,
        expected_headers=expected_headers,
        has_headers=has_headers,
    )
    return CsvImportResult(
        headers=tuple(headers),
        rows=tuple(tuple(row) for row in normalized_rows),
        expected_headers=tuple(expected_headers),
        covariate_names=tuple(covariate_names),
        covariate_types=tuple(covariate_types),
    )


def _validate_headers(
    headers: list[str], expected_headers: list[str] | tuple[str, ...]
) -> None:
    """Require the workspace columns to be present in the expected order."""
    required_count = len(expected_headers)
    if headers[:required_count] == list(expected_headers):
        return
    for index, expected in enumerate(expected_headers):
        found = headers[index] if index < len(headers) else "(missing)"
        if found != expected:
            raise CsvImportError(f'Column {index + 1}: expected "{expected}", found "{found}". Use the template or map this field to a CSV column.')


def _validate_minimum_width(rows: list[list[str]], minimum_width: int) -> None:
    """Reject headerless rows that cannot supply all workspace columns."""
    for row_number, row in enumerate(rows, start=1):
        if len(row) < minimum_width:
            raise CsvImportError(
                f"CSV row {row_number} must contain at least {minimum_width} columns."
            )


def normalize_import_rows(
    rows: list[list[str]] | tuple[tuple[str, ...], ...], *, minimum_width: int = 0
) -> list[list[str]]:
    """Return mutable, rectangular rows for insertion into the Qt model."""
    return tabular_data.normalize_rows(
        [list(row) for row in rows], minimum_width=minimum_width
    )


def _validate_years(rows: list[list[str]], year_column: int) -> None:
    for row_number, row in enumerate(rows, start=1):
        if year_column >= len(row):
            raise CsvImportError(f"The year at row {row_number} is missing.")
        try:
            int(row[year_column])
        except ValueError as exc:
            raise CsvImportError(
                f"The year at row {row_number} is not an integer number."
            ) from exc


def _infer_covariates(
    rows: list[list[str]],
    *,
    headers: list[str],
    expected_headers: list[str] | tuple[str, ...],
    has_headers: bool,
) -> tuple[list[str], list[str]]:
    width = len(rows[0]) if rows else len(headers)
    covariate_count = max(0, width - len(expected_headers))
    if covariate_count == 0:
        return [], []

    names = headers[len(expected_headers) :] if has_headers else []
    names += [""] * (covariate_count - len(names))
    normalized_names = [
        name if name.strip() else f"Covariate {index + 1}"
        for index, name in enumerate(names[:covariate_count])
    ]
    offset = len(expected_headers)
    types = [
        _covariate_type(row[offset + index] for row in rows)
        for index in range(covariate_count)
    ]
    return normalized_names, types


def _covariate_type(values: Iterable[str]) -> str:
    for value in values:
        try:
            float(value)
        except ValueError:
            return "factor"
    return "continuous"


def _xlsx_metadata(path, sheet, selected, row_numbers, columns, notices):
    from zipfile import ZipFile
    from xml.etree import ElementTree
    # Streaming worksheets omit merge and visibility metadata. Inspect the XML
    # first, opening a normal workbook only when that metadata is present.
    with ZipFile(path) as archive:
        xml = archive.read(sheet._worksheet_path)
    root = ElementTree.fromstring(xml)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    metadata = root.find("s:mergeCells", ns) is not None or any(e.get("hidden") == "1" for e in root.iter())
    if metadata:
        _xlsx_visible_metadata(path, selected, row_numbers, columns, notices)


def _xlsx_visible_metadata(path, selected, row_numbers, columns, notices):
    from openpyxl import load_workbook
    normal = load_workbook(path, read_only=False, data_only=False)
    try:
        full = normal[selected]
        for merged in full.merged_cells.ranges:
            if merged.min_row <= row_numbers[-1] and merged.max_row >= row_numbers[0] and merged.min_col <= columns[-1]+1 and merged.max_col >= columns[0]+1:
                raise CsvImportError(f"{selected}!{merged} intersects the data; unmerge these cells first.")
        _xlsx_hidden_notices(full, row_numbers, columns, notices)
    finally:
        normal.close()

def _xlsx_hidden_notices(full, row_numbers, columns, notices):
    hidden_rows = sum(bool(full.row_dimensions[number].hidden) for number in row_numbers)
    hidden_cols = sum(1 for i in columns if any(d.hidden and d.min <= i+1 <= d.max for d in full.column_dimensions.values()))
    if notices is not None and (hidden_rows or hidden_cols):
        for count, label in ((hidden_rows, "row"), (hidden_cols, "column")):
            if count:
                notices.append(f"{count} hidden {label}{' was' if count == 1 else 's were'} imported.")

def _xlsx_cell_value(cell, cached, selected, row_index, output_column, year_column, numeric_columns, dates):
    from datetime import date, datetime
    formula = 0
    value = cell.value
    location = f"{selected}!{cell.coordinate}" if value is not None else ""
    if cell.data_type == "f":
        value = cached.get((cell.row, cell.column))
        if value is None:
            raise CsvImportError(f"{location} has a formula with no saved value. Open and save it in Excel first.")
        formula = 1
    if row_index and output_column in numeric_columns and (cell.data_type == "e" or isinstance(value, str) and value.startswith(("#DIV/0!", "#N/A"))):
        raise CsvImportError(f"{location} contains an Excel error value.")
    if isinstance(value, (datetime, date)):
        if row_index and output_column == year_column:
            dates.append(location)
            value = str(value.year)
        else:
            value = value.isoformat()
    if isinstance(value, float):
        value = str(int(value)) if value.is_integer() else repr(value)
    return "" if value is None else str(value), formula


def _xlsx_select_sheet(book, worksheet):
    names = [sheet.title for sheet in book if sheet.sheet_state == "visible"]
    if not names:
        raise CsvImportError("This workbook has no visible worksheets.")
    if worksheet in names:
        return names, worksheet
    selected = next((name for name in names if _xlsx_has_values(book[name])), names[0])
    return names, selected


def _xlsx_has_values(sheet):
    return any(any(cell.value is not None for cell in row) for row in sheet.iter_rows())


def _xlsx_cells(sheet):
    cells = [list(row) for row in sheet.iter_rows() if any(cell.value is not None for cell in row)]
    if not cells:
        return [], [], []
    row_numbers = [next(cell.row for cell in row if cell.value is not None) for row in cells]
    columns = [i for i in range(max(map(len, cells)))
               if any(i < len(row) and row[i].value is not None for row in cells)]
    return cells, row_numbers, columns


def read_xlsx(path, worksheet=None, *, year_column=None, numeric_columns=(), notices=None):
    """Read stored worksheet cells; preserve coordinates for actionable errors."""
    from openpyxl import load_workbook

    try:
        book = load_workbook(path, read_only=True, data_only=False)
        saved = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CsvImportError("Could not read this workbook; use an unprotected .xlsx file.") from exc
    try:
        names, selected = _xlsx_select_sheet(book, worksheet)
        sheet = book[selected]
        cells, row_numbers, columns = _xlsx_cells(sheet)
        if not cells:
            return [], [], names, selected
        _xlsx_metadata(path, sheet, selected, row_numbers, columns, notices)
        cached = {(c.row, c.column): c.value for row in saved[selected].iter_rows() for c in row if c.value is not None}
        formulas = 0
        dates = []
        result = []
        for row_index, row in enumerate(cells):
            values = []
            for output_column, i in enumerate(columns):
                value, formula = _xlsx_cell_value(
                    row[i], cached, selected, row_index, output_column,
                    year_column, numeric_columns, dates)
                formulas += formula
                values.append(value)
            result.append(values)
        if notices is not None:
            if formulas:
                notices.append(f"{formulas} formula cells use their saved values.")
            if dates:
                notices.append("Dates use their year: " + ", ".join(dates) + ".")
        return result[0], result[1:], names, selected
    finally:
        book.close()
        saved.close()
