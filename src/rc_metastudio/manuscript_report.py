"""Manuscript content taken exclusively from an executed result."""
from dataclasses import dataclass
from html import escape
import io
from PyQt6.QtCore import QRectF, Qt
from PyQt6.QtGui import QImage, QPainter, QPdfWriter, QTextDocument, QPageSize
from PyQt6.QtSvg import QSvgRenderer
from docx import Document
from docx.shared import Mm
from rc_metastudio import __version__


def methods_prose(metric, model, parameters, confidence, workflow, details, fits):
    from rc_metastudio.analysis_method_labels import parameter_value_display_label
    plural = {"odds ratio": "Odds ratios", "risk ratio": "Risk ratios", "hazard ratio": "Hazard ratios",
              "mean difference": "Mean differences", "standardized mean difference": "Standardized mean differences"}.get(metric, metric.capitalize() + " estimates")
    sentence = f"{plural} were pooled using a {model} model"
    estimator = parameters.get('rm.method', parameters.get('estimator'))
    if estimator and model == 'random-effects':
        label = parameter_value_display_label('rm.method', estimator).replace('DerSimonian-Laird', 'DerSimonian–Laird')
        sentence += f" with the {label} estimator of between-study variance"
    parts = [sentence + '.']
    if any('reitsma' in method.lower() for _, method, _ in fits):
        parts = []
    inference = parameters.get('inference.method')
    if inference in ('z', 'normal'):
        parts.append(f"Confidence intervals ({float(confidence):g}%) used the normal approximation.")
    elif inference:
        parts.append(f"Confidence intervals ({float(confidence):g}%) used {parameter_value_display_label('inference.method', inference)}.")
    if workflow == 'subgroup':
        parts.append(f"Subgroup analyses were performed by {parameters.get('cov_name', 'the selected covariate')}.")
        parts.append(f"Subgroup differences were tested using an inverse-variance {model} moderator model.")
    elif workflow == 'meta-regression':
        moderators = next((line.removeprefix('Moderators: ') for line in details.splitlines() if line.startswith('Moderators: ')), '')
        if moderators:
            parts.append(f"Meta-regression was performed with {moderators} as moderators.")
    elif workflow == 'cumulative':
        parts.append('Cumulative meta-analysis was performed in current study order.')
    elif workflow == 'leave-one-out':
        parts.append('Leave-one-out analyses were performed by omitting each study in turn.')
    if any('reitsma' in method.lower() for _, method, _ in fits):
        parts.append('Sensitivity and specificity were jointly modeled using the Reitsma bivariate model.')
    return ' '.join(parts)


@dataclass(frozen=True)
class ManuscriptReport:
    title: str
    methods: str
    sections: tuple
    figures: tuple
    summary: tuple[str, ...]
    comparison: str = ''

    @property
    def results(self):
        return ' '.join(self.summary)

    @classmethod
    def from_result(cls, result):
        from dataclasses import replace
        from rc_metastudio.analysis_results import ResultTable
        methods = result.manuscript_methods
        sections = []
        supplementary = []
        summary = [result.manuscript_context] if result.manuscript_context else []
        merged = []
        for section in result.sections:
            if section.kind != 'text':
                continue
            main = []
            values = {}
            for table in section.tables:
                if table.title.startswith('Calculation scale') or section.source_key == 'Weights':
                    supplementary.append(replace(section, title='Supplementary details', paragraphs=(), tables=(table,)))
                    continue
                main.append(table)
                values = dict(zip(table.headers, table.rows[0])) if len(table.rows) == 1 else {}
                if table.headers == ('Statistic', 'Value'):
                    statistics = dict(table.rows)
                    estimate = statistics.get('Estimate')
                    lower = next((v for k,v in table.rows if k.startswith('Lower bound')), None)
                    upper = next((v for k,v in table.rows if k.startswith('Upper bound')), None)
                    if estimate and lower and upper:
                        interval = next(k.partition('(')[2].rstrip(') ').strip() for k,v in table.rows if k.startswith('Lower bound'))
                        summary.append(f"{table.title} was {estimate} ({interval} {lower} to {upper}).")
                elif len(table.rows) > 1 and 'Estimate' in table.headers and 'Term' in table.headers:
                    for row in table.rows:
                        values = dict(zip(table.headers, row))
                        lower = next((v for k, v in values.items() if k.startswith('Lower bound')), None)
                        upper = next((v for k, v in values.items() if k.startswith('Upper bound')), None)
                        level = next((k.partition('(')[2].rstrip(')') for k in values if k.startswith('Lower bound (')), 'CI')
                        interval = f" ({level} {lower} to {upper})" if lower is not None and upper is not None else ''
                        p = f"; p = {values['p-value']}" if 'p-value' in values else ''
                        summary.append(f"{values['Term']}: coefficient {values['Estimate']}{interval}{p}.")
                if 'Subgroups' in table.headers and 'Estimate' in table.headers:
                    for row in table.rows:
                        cells = dict(zip(table.headers, row))
                        lower = next((v for k, v in cells.items() if k.startswith('Lower bound')), None)
                        upper = next((v for k, v in cells.items() if k.startswith('Upper bound')), None)
                        if lower is not None and upper is not None:
                            level = next(k.partition("(")[2].rstrip(")") for k in cells if k.startswith("Lower bound"))
                            summary.append(f"{cells['Subgroups']}: estimate {cells['Estimate']} ({level} {lower} to {upper}).")
                if table.title == 'Test of subgroup differences' and values:
                    summary.append(f"Between-group test: {values['Test']} = {values['Statistic']} (df = {values['df']}; p = {values['p-value']}).")
                if table.title == 'Residual heterogeneity' and values:
                    summary.append('Residual heterogeneity: ' + ', '.join(f'{key} = {value}' for key, value in values.items()) + '.')
                if section.title == 'Summary' and values:
                    merged.extend((table.title + ': ' + header, value) for header, value in values.items())
                if table.title.endswith('Model Results') and 'Estimate' in values and not any('reitsma' in t.lower() for t in result.texts):
                    lower = next((v for k,v in values.items() if k.startswith('Lower bound')), None)
                    upper = next((v for k,v in values.items() if k.startswith('Upper bound')), None)
                    level = next((k.split('(')[1].rstrip(')') for k in values if k.startswith('Lower bound (')), '95% CI')
                    metric = result.manuscript_title.rsplit('(', 1)[-1].rstrip(')')
                    if lower is not None and upper is not None:
                        p = '; p = ' + values['p-value'] if 'p-value' in values else ''
                        summary.append(f"The pooled {metric} was {values['Estimate']} ({level} {lower} to {upper}{p}).")
                elif table.title.endswith('Heterogeneity') and 'Q' in values:
                    summary.append(f"Heterogeneity: Q = {values['Q']} (df = {values.get('df', '')}; p = {values.get('p-value', '')}), I² = {values.get('I²', '')}" + (f", τ² = {values['τ²']}" if 'τ²' in values else '') + '.')
                elif table.title.endswith('Prediction interval') and len(table.rows) == 1:
                    row = table.rows[0]
                    level = table.headers[0].partition('(')[2].removesuffix(')').replace(' PI', '')
                    summary.append(f"The {level} prediction interval was {row[0]} to {row[1]}.")
            if main and (section.title != 'Summary' or any(t.title == 'Regression coefficients' for t in main)):
                sections.append(replace(section, paragraphs=(), tables=tuple(main)))
        if merged:
            original = next(s for s in result.sections if s.title == 'Summary')
            sections.insert(0, replace(original, title='Summary', paragraphs=(), tables=(ResultTable('Summary', ('Statistic', 'Value'), tuple(merged)),)))
        if any(t.headers[0] == 'Q' for s in result.sections for t in s.tables) and 'Reitsma' not in methods:
            methods += " Heterogeneity was assessed with Cochran's Q and I²."
        methods += f' Analyses were performed in RC MetaStudio {__version__}.'
        packages = ', '.join(p.strip() for p in result.software.split(',') if not p.strip().startswith('RCMetaR'))
        if packages:
            methods += ' R packages loaded for this analysis: ' + packages + '.'
        details = tuple(line for line in result.details.splitlines()
                        if line.startswith(('Order:', 'Each row', 'Subgroups by', 'Moderators:', 'Omission-only fits:')))
        if details:
            supplementary.append(replace(next(s for s in result.sections if s.kind == 'text'),
                                         title='Supplementary details', paragraphs=details, tables=(), value=''))
        figures = tuple((s.title, result.display_images.get(s.source_key, s.value)) for s in result.sections if s.kind == 'image')
        return cls(result.manuscript_title or 'Meta-analysis', methods, tuple(sections + supplementary), figures, tuple(summary), result.manuscript_comparison)

    def plain_text(self):
        parts = [self.title, self.comparison, 'Methods', self.methods, 'Results', self.results]
        number = 0
        for section in self.sections:
            parts.extend(section.paragraphs)
            for table in section.tables:
                number += 1
                parts.append(f'Table {number}. {table.title}')
                parts.extend(('\t'.join(table.headers), *('\t'.join(row) for row in table.rows)))
        return '\n\n'.join(parts)

    def html(self, supplementary=False):
        parts = [f'<h1>{escape(self.title)}</h1><p>{escape(self.comparison)}</p><h2>Methods</h2><p>{escape(self.methods)}</p><h2>Results</h2><p>{escape(self.results)}</p>']
        if supplementary:
            parts = []
        table_number = 0
        for number, section in enumerate(self.sections, 1):
            if (section.title == "Supplementary details") != supplementary:
                continue
            if section.title == 'Supplementary details' and (number == 1 or self.sections[number-2].title != section.title):
                parts.append('<h2>Supplementary details</h2>')
            parts.extend(f'<p>{escape(paragraph)}</p>' for paragraph in section.paragraphs)
            for table in section.tables:
                table_number += 1
                parts.append(f'<p><b>Table {table_number}. {escape(table.title)}</b></p><hr><table border="0" cellspacing="0" cellpadding="5">')
                for index, row in enumerate((table.headers, *table.rows)):
                    tag = 'th' if index == 0 else 'td'
                    parts.append('<tr>' + ''.join(f'<{tag}>' + escape(c) + f'</{tag}>' for c in row) + '</tr>')
                parts.append('</table><hr><p>&nbsp;</p>')
        return ''.join(parts)


def figure_source(path):
    if str(path).lower().endswith('.svg'):
        svg = QSvgRenderer(str(path))
        size = svg.defaultSize()
        return svg, size.width() / size.height()
    image = QImage(str(path))
    if image.isNull():
        raise ValueError(f'Cannot read figure: {path}')
    return image, image.width() / image.height()


def paint_figure(painter, source, rect):
    if isinstance(source, QSvgRenderer):
        source.render(painter, rect)
    else:
        painter.drawImage(rect, source)


def write_docx(report, path):
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.left_margin = section.right_margin = Mm(15)
    document.add_heading(report.title, 0)
    document.add_paragraph(report.comparison)
    document.add_heading('Methods', 1)
    document.add_paragraph(report.methods)
    document.add_heading('Results', 1)
    document.add_paragraph(report.results)
    table_number = 0
    def add_tables(supplementary):
        nonlocal table_number
        for number, section in enumerate(report.sections, 1):
            if (section.title == "Supplementary details") != supplementary:
                continue
            if section.title == 'Supplementary details' and (number == 1 or report.sections[number-2].title != section.title):
                document.add_heading('Supplementary details', 1)
            for paragraph in section.paragraphs or (() if section.tables else (section.value,)):
                document.add_paragraph(paragraph)
            for table in section.tables:
                table_number += 1
                document.add_paragraph(f'Table {table_number}. {table.title}', style='Caption')
                target = document.add_table(rows=1, cols=len(table.headers))
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn
                borders = OxmlElement('w:tblBorders')
                for edge in ('top', 'bottom'):
                    rule = OxmlElement('w:' + edge)
                    rule.set(qn('w:val'), 'single')
                    rule.set(qn('w:sz'), '6')
                    borders.append(rule)
                target._tbl.tblPr.append(borders)
                for row, values in zip([target.rows[0], *[target.add_row() for _ in table.rows]], (table.headers, *table.rows)):
                    for cell, value in zip(row.cells, values):
                        cell.text = value
                document.add_paragraph()
    add_tables(False)
    from PyQt6.QtCore import QBuffer, QIODevice
    for number, (title, path_) in enumerate(report.figures, 1):
        source, ratio = figure_source(path_)
        width = 180 if ratio > 1.4 or 'forest' in title.lower() else 85
        pixel_width = round(width / 25.4 * 300)
        pixel_height = round(pixel_width / ratio)
        page_height = round(230 / 25.4 * 300)
        for top in range(0, pixel_height, page_height):
            image = QImage(pixel_width, min(page_height, pixel_height-top), QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            painter = QPainter(image)
            painter.translate(0, -top)
            paint_figure(painter, source, QRectF(0, 0, pixel_width, pixel_height))
            painter.end()
            buffer = QBuffer()
            buffer.open(QIODevice.OpenModeFlag.WriteOnly)
            image.save(buffer, 'PNG')
            document.add_page_break()
            document.add_picture(io.BytesIO(bytes(buffer.data())), width=Mm(width))
            document.add_paragraph(f'Figure {number}. {title}' + (' (continued)' if top else ''))
    add_tables(True)
    document.save(path)


def write_pdf(report, path):
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setResolution(600)
    document = QTextDocument()
    document.setHtml(report.html())
    # QTextDocument prints paginated real tables; figures get their own pages so
    # long reports cannot overwrite a figure or split its caption.
    from PyQt6.QtCore import QSizeF
    from PyQt6.QtGui import QAbstractTextDocumentLayout
    scale = writer.resolution() / 96
    page = writer.pageLayout().paintRectPixels(writer.resolution())
    width, height = page.width() / scale, page.height() / scale
    document.setPageSize(QSizeF(width, height))
    painter = QPainter(writer)
    def print_pages(first_page=False):
        for number in range(document.pageCount()):
            if number or first_page:
                writer.newPage()
            painter.save()
            painter.scale(scale, scale)
            painter.setClipRect(QRectF(0, 0, width, height))
            painter.translate(0, -number * height)
            context = QAbstractTextDocumentLayout.PaintContext()
            context.clip = QRectF(0, number * height, width, height)
            document.documentLayout().draw(painter, context)
            painter.restore()
    print_pages()
    for number, (title, path_) in enumerate(report.figures, 1):
        writer.newPage()
        source, ratio = figure_source(path_)
        mm = 180 if ratio > 1.4 or 'forest' in title.lower() else 85
        w = mm / 25.4 * writer.resolution()
        h = w / ratio
        caption = QTextDocument()
        caption.setTextWidth(width)
        caption.setHtml(f'<b>Figure {number}. {escape(title)} (continued)</b>')
        # Lay out captions in the same 96-DPI coordinates as table captions.
        available = page.height() - (caption.size().height() + 12) * scale
        import math
        for part in range(math.ceil(h / available)):
            if part:
                writer.newPage()
            painter.save()
            painter.setClipRect(QRectF(0, 0, page.width(), available))
            rect = QRectF((page.width() - w) / 2, -part * available, w, h)
            paint_figure(painter, source, rect)
            painter.restore()
            painter.save()
            painter.scale(scale, scale)
            caption_y = min(h - part * available, available) / scale + 12
            caption.setHtml(f'<b>Figure {number}. {escape(title)}' + (' (continued)' if part else '') + '</b>')
            painter.translate(0, caption_y)
            caption.drawContents(painter)
            painter.restore()
    if any(section.title == 'Supplementary details' for section in report.sections):
        document.setHtml(report.html(supplementary=True))
        print_pages(first_page=True)
    painter.end()
